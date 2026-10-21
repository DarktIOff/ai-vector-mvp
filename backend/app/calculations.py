from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .catalog import (
    METRIC_BY_CODE,
    METRIC_BY_COLUMN,
    resolve_header_metrics,
)
from .workbooks import (
    decimal_value,
    find_sheet,
    normalize_text,
    open_reader,
    recognize_header,
)


BILLION_FROM_MSP = Decimal("1000000")
BILLION_FROM_RF = Decimal("1000")
REPORT_TAX_COLUMNS = (6, 8, 10, 30, 31, 32, 33, 34, 35)


@dataclass(frozen=True)
class Evidence:
    file: str
    absolute_path: str
    sheet: str
    cell: str
    row_code: int | None
    report_column: int | None
    report_column_name: str
    section: str | None
    territory: str
    period_end: str
    raw_value: str
    raw_unit: str
    conversion: str | None = None
    formula: str | None = None
    lineage: tuple[dict, ...] = ()

    def to_dict(self) -> dict:
        result = self.__dict__.copy()
        result["lineage"] = list(self.lineage)
        return result


@dataclass(frozen=True)
class ValueWithEvidence:
    value: Decimal
    evidence: Evidence


def safe_ratio(numerator: Decimal, denominator: Decimal) -> Decimal | None:
    if denominator == 0:
        return None
    return numerator / denominator * Decimal("100")


def _find_code_row(sheet, code_column: int, wanted: int) -> int:
    for row in range(1, sheet.max_row + 1):
        value = decimal_value(sheet.value(row, code_column))
        if value is not None and value == Decimal(wanted):
            return row
    raise ValueError(
        f"Не найден код строки {wanted} на листе «{sheet.title}»"
    )


class MspWorkbookSnapshot:
    """One-pass, in-memory view of the semantic facts used by a report."""

    def __init__(self, path: Path, period_end: str):
        self.path = path
        self.period_end = period_end
        self._values: dict[tuple[str, int, int], ValueWithEvidence] = {}
        reader = open_reader(path)
        try:
            for section, offset in (("I", 0), ("II", 1000)):
                sheet_name = find_sheet(reader, f"Раздел {section}")
                if sheet_name is None:
                    raise ValueError(f"В {path.name} не найден Раздел {section}")
                sheet = reader.sheet(sheet_name)
                schema = recognize_header(sheet)
                resolution = resolve_header_metrics(schema.raw_headers)
                if not resolution.is_complete:
                    raise ValueError(
                        "Не удалось семантически сопоставить заголовки "
                        f"{path.name}: {sorted(resolution.unmapped)}"
                    )
                for row in range(schema.header_number_row + 1, sheet.max_row + 1):
                    raw_code = decimal_value(sheet.value(row, schema.code_column))
                    if (
                        raw_code is None
                        or raw_code != raw_code.to_integral_value()
                    ):
                        continue
                    source_code = int(raw_code)
                    base_code = source_code - offset
                    if not 1000 <= base_code < 2000:
                        continue
                    for source_column, physical_column in schema.metric_columns.items():
                        metric_code = resolution.source_to_code.get(
                            source_column
                        )
                        metric = METRIC_BY_CODE.get(metric_code or "")
                        if metric is None:
                            raise ValueError(
                                f"Нет показателя для графы {source_column}: "
                                f"{schema.raw_headers.get(source_column, '')}"
                            )
                        report_column = metric.column
                        value = decimal_value(sheet.value(row, physical_column))
                        if value is None:
                            continue
                        evidence = Evidence(
                            file=path.name,
                            absolute_path=str(path.resolve()),
                            sheet=sheet_name,
                            cell=sheet.coordinate(row, physical_column),
                            row_code=source_code,
                            report_column=report_column,
                            report_column_name=METRIC_BY_COLUMN[report_column].name,
                            section=section,
                            territory="Российская Федерация",
                            period_end=period_end,
                            raw_value=str(value),
                            raw_unit=(
                                "тыс. руб." if report_column != 1 else "ед."
                            ),
                            conversion=(
                                "тыс. руб. ÷ 1 000 000 = млрд руб."
                                if report_column != 1
                                else None
                            ),
                            lineage=(
                                {
                                    "semantic_code": metric.code,
                                    "source_report_column": source_column,
                                    "source_header": schema.raw_headers.get(
                                        source_column,
                                        "",
                                    ),
                                    "mapping_method": "normalized_header",
                                },
                            ),
                        )
                        self._values[(section, base_code, report_column)] = (
                            ValueWithEvidence(value, evidence)
                        )
        finally:
            reader.close()

    def value(
        self, section: str, base_code: int, report_column: int
    ) -> ValueWithEvidence:
        if section == "III":
            section_i = self.value("I", base_code, report_column)
            section_ii = self.value("II", base_code, report_column)
            value = section_i.value - section_ii.value
            return ValueWithEvidence(
                value,
                Evidence(
                    file=self.path.name,
                    absolute_path=str(self.path.resolve()),
                    sheet="Производный показатель",
                    cell="—",
                    row_code=base_code,
                    report_column=report_column,
                    report_column_name=METRIC_BY_COLUMN[report_column].name,
                    section="III",
                    territory="Российская Федерация",
                    period_end=self.period_end,
                    raw_value=str(value),
                    raw_unit="тыс. руб.",
                    formula="Раздел III = Раздел I − Раздел II",
                    lineage=(
                        section_i.evidence.to_dict(),
                        section_ii.evidence.to_dict(),
                    ),
                ),
            )
        try:
            return self._values[(section, base_code, report_column)]
        except KeyError as exc:
            raise ValueError(
                f"Нет значения: {self.path.name}, Раздел {section}, "
                f"строка {base_code}, графа {report_column}"
            ) from exc


def msp_value(
    path: Path,
    section: str,
    base_code: int,
    report_column: int,
    period_end: str,
) -> ValueWithEvidence:
    return MspWorkbookSnapshot(path, period_end).value(
        section, base_code, report_column
    )


def _find_label_row(sheet, wanted: str) -> int:
    wanted_normalized = normalize_text(wanted)
    for row in range(1, sheet.max_row + 1):
        if wanted_normalized in normalize_text(sheet.value(row, 1)):
            return row
    raise ValueError(f"Не найдена строка «{wanted}»")


def national_receipts(
    path: Path, old_year: int, new_year: int
) -> tuple[ValueWithEvidence, ValueWithEvidence, ValueWithEvidence, ValueWithEvidence]:
    reader = open_reader(path)
    try:
        sheet = reader.sheet(reader.sheet_names[0])
        old_column = new_column = 0
        for row in range(1, min(sheet.max_row, 20) + 1):
            for column in range(1, sheet.max_column + 1):
                if normalize_text(sheet.value(row, column)) != "поступило":
                    continue
                for header_row in range(row + 1, min(row + 5, sheet.max_row) + 1):
                    for candidate in range(column, min(column + 8, sheet.max_column) + 1):
                        text = str(sheet.value(header_row, candidate) or "")
                        if str(old_year) in text:
                            old_column = candidate
                        if str(new_year) in text:
                            new_column = candidate
                    if old_column and new_column:
                        break
        if not old_column or not new_column:
            raise ValueError("Не определены колонки блока «ПОСТУПИЛО»")
        budget_row = _find_label_row(
            sheet, "Всего доходов консолидированный бюджет РФ"
        )
        insurance_row = _find_label_row(
            sheet, "Всего доходов по страховым взносам"
        )

        def read(row: int, column: int, year: int, label: str) -> ValueWithEvidence:
            value = decimal_value(sheet.value(row, column))
            if value is None:
                raise ValueError(
                    f"Ожидалось число в {path.name}!{sheet.coordinate(row, column)}"
                )
            evidence = Evidence(
                file=path.name,
                absolute_path=str(path.resolve()),
                sheet=sheet.title,
                cell=sheet.coordinate(row, column),
                row_code=None,
                report_column=None,
                report_column_name=label,
                section=None,
                territory="Российская Федерация",
                period_end=f"{year}-06-30",
                raw_value=str(value),
                raw_unit="млн руб.",
                conversion="млн руб. ÷ 1 000 = млрд руб.",
            )
            return ValueWithEvidence(value, evidence)

        return (
            read(budget_row, old_column, old_year, "Доходы консолидированного бюджета"),
            read(budget_row, new_column, new_year, "Доходы консолидированного бюджета"),
            read(insurance_row, old_column, old_year, "Страховые взносы по РФ"),
            read(insurance_row, new_column, new_year, "Страховые взносы по РФ"),
        )
    finally:
        reader.close()


def _metric_pair(
    old_snapshot: MspWorkbookSnapshot,
    new_snapshot: MspWorkbookSnapshot,
    section: str,
    base_code: int,
    report_column: int,
) -> dict:
    old = old_snapshot.value(section, base_code, report_column)
    new = new_snapshot.value(section, base_code, report_column)
    old_billion = old.value / BILLION_FROM_MSP
    new_billion = new.value / BILLION_FROM_MSP
    return {
        "old": old_billion,
        "new": new_billion,
        "growth_rate": safe_ratio(new_billion, old_billion),
        "difference": new_billion - old_billion,
        "sources": [old.evidence.to_dict(), new.evidence.to_dict()],
    }


def _industry_definitions(codes_path: Path) -> list[dict]:
    workbook = load_workbook(codes_path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        result: list[dict] = []
        for row in sheet.iter_rows(values_only=True):
            if len(row) < 4:
                continue
            code = decimal_value(row[2])
            section = normalize_text(row[3])
            if (
                code is None
                or code != code.to_integral_value()
                or section != normalize_text("Раздел I")
                or int(code) == 1010
            ):
                continue
            result.append(
                {
                    "name": clean_industry_name(str(row[0] or "")),
                    "okved": str(row[1] or ""),
                    "base_row_code": int(code),
                }
            )
        return result
    finally:
        workbook.close()


def clean_industry_name(value: str) -> str:
    text = value.split("(стр.")[0].split("(стр ")[0].strip()
    return text.rstrip("–- ").strip()


def find_receipts_workbook(root: Path) -> Path:
    candidates = [
        path
        for path in root.glob("*.xls*")
        if "поступлен" in normalize_text(path.name)
        and "1-мсп" not in normalize_text(path.name)
    ]
    if not candidates:
        raise FileNotFoundError("Не найдена книга общероссийских поступлений")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def build_report_data(
    old_path: Path,
    new_path: Path,
    receipts_path: Path,
    codes_path: Path,
    *,
    old_year: int = 2025,
    new_year: int = 2026,
    months_count: int = 6,
) -> dict:
    old_period = f"{old_year}-{months_count:02d}-30"
    new_period = f"{new_year}-{months_count:02d}-30"
    old_snapshot = MspWorkbookSnapshot(old_path, old_period)
    new_snapshot = MspWorkbookSnapshot(new_path, new_period)
    rf_old_budget, rf_new_budget, rf_old_insurance, rf_new_insurance = (
        national_receipts(receipts_path, old_year, new_year)
    )

    def billion(value: ValueWithEvidence, divisor: Decimal) -> Decimal:
        return value.value / divisor

    rf_old_total = billion(rf_old_budget, BILLION_FROM_RF) + billion(
        rf_old_insurance, BILLION_FROM_RF
    )
    rf_new_total = billion(rf_new_budget, BILLION_FROM_RF) + billion(
        rf_new_insurance, BILLION_FROM_RF
    )
    msp_budget = _metric_pair(
        old_snapshot, new_snapshot, "I", 1010, 4
    )
    msp_insurance = _metric_pair(
        old_snapshot, new_snapshot, "I", 1010, 36
    )
    micro_budget = _metric_pair(
        old_snapshot, new_snapshot, "II", 1010, 4
    )
    micro_insurance = _metric_pair(
        old_snapshot, new_snapshot, "II", 1010, 36
    )

    def pair(old: Decimal, new: Decimal, sources: list[dict]) -> dict:
        return {
            "old": old,
            "new": new,
            "growth_rate": safe_ratio(new, old),
            "difference": new - old,
            "sources": sources,
        }

    msp_total = pair(
        msp_budget["old"] + msp_insurance["old"],
        msp_budget["new"] + msp_insurance["new"],
        [*msp_budget["sources"], *msp_insurance["sources"]],
    )
    micro_total = pair(
        micro_budget["old"] + micro_insurance["old"],
        micro_budget["new"] + micro_insurance["new"],
        [*micro_budget["sources"], *micro_insurance["sources"]],
    )
    small_medium_total = pair(
        msp_total["old"] - micro_total["old"],
        msp_total["new"] - micro_total["new"],
        [*msp_total["sources"], *micro_total["sources"]],
    )
    rf_total = pair(
        rf_old_total,
        rf_new_total,
        [
            rf_old_budget.evidence.to_dict(),
            rf_old_insurance.evidence.to_dict(),
            rf_new_budget.evidence.to_dict(),
            rf_new_insurance.evidence.to_dict(),
        ],
    )
    msp_share = pair(
        safe_ratio(msp_total["old"], rf_old_total) or Decimal("0"),
        safe_ratio(msp_total["new"], rf_new_total) or Decimal("0"),
        [*msp_total["sources"], *rf_total["sources"]],
    )
    table1 = [
        {"code": "rf_total", "name": "Поступило по РФ", **rf_total},
        {
            "code": "rf_budget",
            "name": "В консолидированный бюджет РФ",
            **pair(
                billion(rf_old_budget, BILLION_FROM_RF),
                billion(rf_new_budget, BILLION_FROM_RF),
                [rf_old_budget.evidence.to_dict(), rf_new_budget.evidence.to_dict()],
            ),
        },
        {
            "code": "rf_insurance",
            "name": "Страховые взносы по РФ",
            **pair(
                billion(rf_old_insurance, BILLION_FROM_RF),
                billion(rf_new_insurance, BILLION_FROM_RF),
                [
                    rf_old_insurance.evidence.to_dict(),
                    rf_new_insurance.evidence.to_dict(),
                ],
            ),
        },
        {"code": "msp_total", "name": "Поступило от МСП", **msp_total},
        {
            "code": "msp_share",
            "name": "Доля МСП в РФ, %",
            **msp_share,
            "growth_rate": None,
        },
        {"code": "msp_budget", "name": "МСП: консолидированный бюджет", **msp_budget},
        {
            "code": "msp_insurance",
            "name": "МСП: страховые взносы",
            **msp_insurance,
        },
        {"code": "micro_total", "name": "Микропредприятия", **micro_total},
        {"code": "micro_budget", "name": "Микро: бюджет", **micro_budget},
        {"code": "micro_insurance", "name": "Микро: страховые взносы", **micro_insurance},
        {
            "code": "small_medium_total",
            "name": "Малые и средние без микро",
            **small_medium_total,
        },
        {
            "code": "small_medium_budget",
            "name": "Малые и средние: бюджет",
            **pair(
                msp_budget["old"] - micro_budget["old"],
                msp_budget["new"] - micro_budget["new"],
                [*msp_budget["sources"], *micro_budget["sources"]],
            ),
        },
        {
            "code": "small_medium_insurance",
            "name": "Малые и средние: страховые взносы",
            **pair(
                msp_insurance["old"] - micro_insurance["old"],
                msp_insurance["new"] - micro_insurance["new"],
                [*msp_insurance["sources"], *micro_insurance["sources"]],
            ),
        },
    ]

    table2 = [
        {
            "code": "msp_budget",
            "name": "Поступило от МСП в консолидированный бюджет",
            "group": "all_msp",
            **msp_budget,
        }
    ]
    for report_column in REPORT_TAX_COLUMNS:
        for section, group in (
            ("I", "all_msp"),
            ("II", "micro"),
            ("III", "small_and_medium_without_micro"),
        ):
            values = _metric_pair(
                old_snapshot,
                new_snapshot,
                section,
                1010,
                report_column,
            )
            table2.append(
                {
                    "code": METRIC_BY_COLUMN[report_column].code,
                    "name": METRIC_BY_COLUMN[report_column].name,
                    "group": group,
                    **values,
                }
            )

    industries: list[dict] = []
    total_current = msp_budget["new"]
    for definition in _industry_definitions(codes_path):
        groups: dict[str, dict] = {}
        for section, group in (
            ("I", "all_msp"),
            ("II", "micro"),
            ("III", "small_and_medium_without_micro"),
        ):
            groups[group] = _metric_pair(
                old_snapshot,
                new_snapshot,
                section,
                definition["base_row_code"],
                4,
            )
        industries.append(
            {
                **definition,
                "groups": groups,
                "growth_rate": groups["all_msp"]["growth_rate"],
                "difference": groups["all_msp"]["difference"],
                "share": safe_ratio(groups["all_msp"]["new"], total_current),
            }
        )
    industries.sort(
        key=lambda item: item["growth_rate"]
        if item["growth_rate"] is not None
        else Decimal("-1e99"),
        reverse=True,
    )

    return {
        "report_type": "1-msp",
        "old_year": old_year,
        "new_year": new_year,
        "months_count": months_count,
        "period_text": period_words(months_count, new_year),
        "table1": table1,
        "table2": table2,
        "industries": industries,
        "warnings": [],
        "input_files": {
            "old_msp": str(old_path.resolve()),
            "new_msp": str(new_path.resolve()),
            "national_receipts": str(receipts_path.resolve()),
            "codes": str(codes_path.resolve()),
        },
    }


MONTH_HEADERS = {
    1: "ЯНВ",
    2: "ЯНВ–ФЕВ",
    3: "ЯНВ–МАРТ",
    4: "ЯНВ–АПР",
    5: "ЯНВ–МАЙ",
    6: "ЯНВ–ИЮНЬ",
    7: "ЯНВ–ИЮЛЬ",
    8: "ЯНВ–АВГ",
    9: "ЯНВ–СЕНТ",
    10: "ЯНВ–ОКТ",
    11: "ЯНВ–НОЯБ",
    12: "ЯНВ–ДЕКАБРЬ",
}

MONTH_WORDS = {
    1: "январь",
    2: "январь–февраль",
    3: "январь–март",
    4: "январь–апрель",
    5: "январь–май",
    6: "январь–июнь",
    7: "январь–июль",
    8: "январь–август",
    9: "январь–сентябрь",
    10: "январь–октябрь",
    11: "январь–ноябрь",
    12: "январь–декабрь",
}


def period_header(months_count: int, year: int) -> str:
    return f"{MONTH_HEADERS[months_count]} {year}"


def period_words(months_count: int, year: int) -> str:
    return f"{MONTH_WORDS[months_count]} {year} года"


def ru_number(
    value: Decimal | None, digits: int = 1, show_plus: bool = False
) -> str:
    if value is None:
        return "—"
    quantum = Decimal("1").scaleb(-digits)
    rounded = value.quantize(quantum, rounding=ROUND_HALF_UP)
    if rounded == 0:
        # Decimal preserves the sign of a rounded zero (for example -0.04 ->
        # Decimal('-0.0')). It is not a meaningful value in a report.
        rounded = abs(rounded)
    sign = "+" if show_plus and rounded > 0 else ""
    formatted = f"{abs(rounded):,.{digits}f}" if rounded < 0 else f"{rounded:,.{digits}f}"
    if rounded < 0:
        formatted = "-" + f"{abs(rounded):,.{digits}f}"
    return sign + formatted.replace(",", "\u00a0").replace(".", ",")


def json_ready(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dict):
        return {key: json_ready(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_ready(item) for item in value]
    if isinstance(value, tuple):
        return [json_ready(item) for item in value]
    return value
