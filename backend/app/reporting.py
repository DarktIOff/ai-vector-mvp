from __future__ import annotations

import copy
import json
import os
import re
import uuid
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from .calculations import (
    build_report_data,
    json_ready,
    period_header,
    ru_number,
)


TAG_RE = re.compile(r"\{[^{}]+\}")
NUMBER_RE = re.compile(
    r"(?<![\w])[-+]?\d+(?:[ \u00a0]\d{3})*(?:[,.]\d+)?%?"
)
REPORT_PRECISION_LEVELS = (1, 3, 6, 9)


def _report_row_digits(*values: Decimal | None) -> int:
    """Keep small report values and a real old/new change visible."""
    present = [value for value in values if value is not None]
    if not present:
        return REPORT_PRECISION_LEVELS[0]
    for digits in REPORT_PRECISION_LEVELS:
        quantum = Decimal("1").scaleb(-digits)
        rounded = [
            value.quantize(quantum, rounding=ROUND_HALF_UP)
            for value in present
        ]
        if any(
            value != 0 and result == 0
            for value, result in zip(present, rounded)
        ):
            continue
        if (
            len(present) >= 2
            and present[0] != present[1]
            and rounded[0] == rounded[1]
        ):
            continue
        return digits
    return REPORT_PRECISION_LEVELS[-1]


def _report_number(
    value: Decimal | None,
    *row_values: Decimal | None,
    show_plus: bool = False,
    min_digits: int = 0,
) -> str:
    context = row_values or (value,)
    return ru_number(
        value,
        digits=max(min_digits, _report_row_digits(*context)),
        show_plus=show_plus,
    )


def _cell_text(cell) -> str:
    return "\n".join(paragraph.text for paragraph in cell.paragraphs)


def _set_cell(cell, value: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph(value)
        return
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = value
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def _table_contains(table: Table, marker: str) -> bool:
    return any(marker in _cell_text(cell) for row in table.rows for cell in row.cells)


def _marker_column(table: Table, marker: str) -> int:
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if marker in _cell_text(cell):
                return index
    raise ValueError(f"В шаблоне не найден маркер {marker}")


def _marker_row(table: Table, marker: str) -> int:
    for index, row in enumerate(table.rows):
        if any(marker in _cell_text(cell) for cell in row.cells):
            return index
    raise ValueError(f"В шаблоне не найден маркер {marker}")


def _clean_marker(row, marker: str) -> None:
    for cell in row.cells:
        if marker in _cell_text(cell):
            _set_cell(cell, _cell_text(cell).replace(marker, "").strip())


def _delete_table_column(table: Table, column_index: int) -> None:
    """Physically remove a service column from the DOCX table XML."""
    grid_columns = table._tbl.tblGrid.gridCol_lst
    if column_index < len(grid_columns):
        table._tbl.tblGrid.remove(grid_columns[column_index])
    for row in table.rows:
        cells = row._tr.tc_lst
        if column_index < len(cells):
            row._tr.remove(cells[column_index])


def _fill_table1(
    table: Table,
    data: list[dict],
    old_year: int,
    new_year: int,
    months: int,
    *,
    use_data_labels: bool = False,
) -> None:
    columns = [_marker_column(table, f"{{tab1_col_{index}}}") for index in range(1, 5)]
    headers = [
        period_header(months, old_year),
        period_header(months, new_year),
        (
            f"Темп {new_year} к {old_year}, %"
            if use_data_labels
            else "Темп, %"
        ),
        "Изменение" if use_data_labels else "Откл.",
    ]
    for column, text in zip(columns, headers, strict=True):
        _set_cell(table.rows[0].cells[column], text)

    def fill_values(row, item: dict) -> None:
        min_digits = (
            3
            if use_data_labels and not item["code"].startswith("rf_")
            else 0
        )
        row_values = (item["old"], item["new"], item["difference"])
        if item["code"] in {"msp_share", "region_share_rf_msp"}:
            values = [
                _report_number(
                    item["old"], *row_values, min_digits=min_digits
                )
                + "%",
                _report_number(
                    item["new"], *row_values, min_digits=min_digits
                )
                + "%",
                "х",
                _report_number(
                    item["difference"],
                    *row_values,
                    show_plus=True,
                    min_digits=min_digits,
                )
                + " п.п.",
            ]
        else:
            values = [
                _report_number(
                    item["old"], *row_values, min_digits=min_digits
                ),
                _report_number(
                    item["new"], *row_values, min_digits=min_digits
                ),
                ru_number(item["growth_rate"]),
                _report_number(
                    item["difference"],
                    *row_values,
                    show_plus=True,
                    min_digits=min_digits,
                ),
            ]
        for column, value in zip(columns, values, strict=True):
            _set_cell(row.cells[column], value)

    if use_data_labels and len(table.rows) == 6:
        items_by_code = {item["code"]: item for item in data}
        summary_codes = (
            "rf_msp_total",
            "region_msp_total",
            "region_msp_budget",
            "region_msp_insurance",
            "region_share_rf_msp",
        )
        for row, code in zip(table.rows[1:], summary_codes, strict=True):
            item = items_by_code[code]
            _set_cell(row.cells[0], item["name"])
            fill_values(row, item)
        return

    for marker_index, item in enumerate(data, start=1):
        marker = f"{{tab1_row{marker_index}}}"
        try:
            row_index = _marker_row(table, marker)
        except ValueError:
            continue
        row = table.rows[row_index]
        if use_data_labels:
            marker_cell = next(
                cell for cell in row.cells if marker in _cell_text(cell)
            )
            _set_cell(marker_cell, item["name"])
        else:
            _clean_marker(row, marker)
        fill_values(row, item)


def _fill_table2(
    table: Table,
    data: list[dict],
    old_year: int,
    new_year: int,
    months: int,
    *,
    regional: bool = False,
) -> None:
    marker_column = _marker_column(table, "{tab2_row1}")
    columns = [_marker_column(table, f"{{tab2_col_{index}}}") for index in range(1, 5)]
    headers = [
        period_header(months, old_year),
        period_header(months, new_year),
        f"Темп {new_year} к {old_year}, %" if regional else "Темп, %",
        "Изменение" if regional else "Откл.",
    ]
    for column, text in zip(columns, headers, strict=True):
        _set_cell(table.rows[0].cells[column], text)
    for marker_index, item in enumerate(data, start=1):
        marker = f"{{tab2_row{marker_index}}}"
        row_index = _marker_row(table, marker)
        row = table.rows[row_index]
        _clean_marker(row, marker)
        if regional and marker_index == 1:
            _set_cell(row.cells[0], item["name"])
        row_values = (item["old"], item["new"], item["difference"])
        values = [
            _report_number(
                item["old"], *row_values, min_digits=3 if regional else 0
            ),
            _report_number(
                item["new"], *row_values, min_digits=3 if regional else 0
            ),
            ru_number(item["growth_rate"]),
            _report_number(
                item["difference"],
                *row_values,
                min_digits=3 if regional else 0,
            ),
        ]
        for column, value in zip(columns, values, strict=True):
            _set_cell(row.cells[column], value)
    _delete_table_column(table, marker_column)


def _resize_table(table: Table, wanted_rows: int) -> None:
    while len(table.rows) > wanted_rows:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < wanted_rows:
        template_row = copy.deepcopy(table.rows[-1]._tr)
        table._tbl.append(template_row)


def _fill_table3(
    table: Table,
    report_data: dict,
    old_year: int,
    new_year: int,
    months: int,
) -> None:
    columns = [_marker_column(table, f"{{tab3_col{index}}}") for index in range(1, 7)]
    _set_cell(table.rows[0].cells[columns[1]], period_header(months, old_year))
    _set_cell(table.rows[0].cells[columns[2]], period_header(months, new_year))
    is_regional = report_data.get("report_mode") == "region_vs_rf"
    if is_regional:
        _set_cell(table.rows[0].cells[columns[3]], "Изменение")
        _set_cell(
            table.rows[0].cells[columns[4]],
            f"Темп {new_year} к {old_year}, %",
        )
        _set_cell(
            table.rows[0].cells[columns[5]],
            "Доля в бюджете региона, %",
        )
    for row in table.rows:
        for cell in row.cells:
            cleaned = _cell_text(cell)
            for index in range(1, 7):
                cleaned = cleaned.replace(f"{{tab3_col{index}}}", "")
            _set_cell(cell, cleaned.strip())

    industries = report_data["industries"]
    wanted_rows = 3 + len(industries) * 3
    _resize_table(table, wanted_rows)
    budget_code = (
        "region_msp_budget"
        if report_data.get("report_mode") == "region_vs_rf"
        else "msp_budget"
    )
    msp_budget = next(
        item for item in report_data["table1"] if item["code"] == budget_code
    )

    def fill(
        row_index: int,
        name: str,
        okved: str,
        values: dict,
        share: Decimal | None,
    ) -> None:
        row_values = (values["old"], values["new"], values["difference"])
        min_digits = 3 if is_regional else 0
        output = [
            name,
            okved,
            _report_number(
                values["old"], *row_values, min_digits=min_digits
            ),
            _report_number(
                values["new"], *row_values, min_digits=min_digits
            ),
            _report_number(
                values["difference"], *row_values, min_digits=min_digits
            ),
            ru_number(values["growth_rate"]),
            _report_number(share, min_digits=min_digits),
        ]
        for column, text in enumerate(output):
            _set_cell(table.rows[row_index].cells[column], text)

    fill(
        2,
        (
            "Поступления от МСП в консолидированный бюджет региона "
            f"«{report_data['territory']}», всего (без страховых взносов)"
            if is_regional
            else "Поступления в консолидированный бюджет РФ"
        ),
        "Х",
        msp_budget,
        None,
    )
    row_index = 3
    for industry in industries:
        for group, label in (
            ("all_msp", industry["name"]),
            ("micro", "микро"),
            ("small_and_medium_without_micro", "малое и среднее"),
        ):
            values = industry["groups"][group]
            share = (
                values["new"] / msp_budget["new"] * Decimal("100")
                if msp_budget["new"] != 0
                else None
            )
            fill(
                row_index,
                label,
                industry["okved"] if group == "all_msp" else "",
                values,
                share,
            )
            row_index += 1


def deterministic_narrative(data: dict) -> str:
    if data.get("report_mode") == "region_vs_rf":
        region = next(
            item for item in data["table1"] if item["code"] == "region_msp_total"
        )
        rf = next(
            item for item in data["table1"] if item["code"] == "rf_msp_total"
        )
        share = next(
            item
            for item in data["table1"]
            if item["code"] == "region_share_rf_msp"
        )
        region_budget = next(
            item
            for item in data["table1"]
            if item["code"] == "region_msp_budget"
        )
        region_insurance = next(
            item
            for item in data["table1"]
            if item["code"] == "region_msp_insurance"
        )
        direction = "выросли" if region["difference"] >= 0 else "снизились"
        rf_direction = "выросли" if rf["difference"] >= 0 else "снизились"
        share_direction = (
            "увеличилась" if share["difference"] >= 0 else "снизилась"
        )
        rf_values = (rf["old"], rf["new"], rf["difference"])
        share_values = (share["old"], share["new"], share["difference"])
        pace_gap = region["growth_rate"] - rf["growth_rate"]
        pace_comparison = (
            "Темп роста региона был выше темпа по РФ на "
            if pace_gap >= 0
            else "Темп роста региона был ниже темпа по РФ на "
        )
        return (
            f"За {data['period_text']} совокупные поступления от МСП "
            "(поступления в консолидированный бюджет и страховые взносы) "
            "в регионе "
            f"«{data['territory']}» {direction} на "
            f"{ru_number(abs(region['growth_rate'] - Decimal('100')))}% и "
            f"составили {ru_number(region['new'], digits=3)} млрд рублей: "
            f"{ru_number(region_budget['new'], digits=3)} млрд рублей "
            "в консолидированный бюджет плюс "
            f"{ru_number(region_insurance['new'], digits=3)} млрд рублей "
            "страховых взносов. В предыдущем периоде соответственно: "
            f"{ru_number(region_budget['old'], digits=3)} + "
            f"{ru_number(region_insurance['old'], digits=3)} = "
            f"{ru_number(region['old'], digits=3)} млрд рублей. "
            "По Российской Федерации совокупные поступления от МСП "
            f"{rf_direction} на "
            f"{ru_number(abs(rf['growth_rate'] - Decimal('100')))}% и "
            "составили "
            f"{_report_number(rf['new'], *rf_values)} млрд рублей. "
            f"{pace_comparison}{ru_number(abs(pace_gap))} п.п. "
            "Доля региона в совокупных поступлениях от МСП по РФ "
            f"{share_direction} с "
            f"{_report_number(share['old'], *share_values, min_digits=3)}% "
            f"до {_report_number(share['new'], *share_values, min_digits=3)}%. "
            "Налоговая и отраслевая детализации далее показывают только "
            "поступления в консолидированный бюджет, без страховых взносов. "
            "Все значения рассчитаны детерминированно; причинные объяснения "
            "без подтверждённого источника не добавлялись."
        )
    msp = next(item for item in data["table1"] if item["code"] == "msp_total")
    rf = next(item for item in data["table1"] if item["code"] == "rf_total")
    share = next(item for item in data["table1"] if item["code"] == "msp_share")
    direction = "выросли" if msp["difference"] >= 0 else "снизились"
    budget_direction = "выросли" if rf["difference"] >= 0 else "снизились"
    share_direction = "увеличилась" if share["difference"] >= 0 else "снизилась"
    msp_values = (msp["old"], msp["new"], msp["difference"])
    rf_values = (rf["old"], rf["new"], rf["difference"])
    share_values = (share["old"], share["new"], share["difference"])
    return (
        f"За {data['period_text']} поступления от МСП {direction} на "
        f"{ru_number(abs(msp['growth_rate'] - Decimal('100')))}% и составили "
        f"{_report_number(msp['new'], *msp_values)} млрд рублей. "
        f"Совокупные доходы бюджетной системы {budget_direction} до "
        f"{_report_number(rf['new'], *rf_values)} млрд рублей. "
        f"Доля МСП в общих поступлениях страны {share_direction} с "
        f"{_report_number(share['old'], *share_values)}% до "
        f"{_report_number(share['new'], *share_values)}%. "
        "Все значения рассчитаны детерминированно; причинные объяснения без "
        "подтверждённого источника не добавлялись."
    )


def _replace_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def _fill_narrative(document: Document, data: dict) -> None:
    if len(document.paragraphs) > 1:
        _replace_paragraph(document.paragraphs[1], deterministic_narrative(data))
    tax_rows = [
        item
        for item in data["table2"]
        if item["group"] == "all_msp"
        and item["code"] not in {"msp_budget", "region_msp_budget"}
    ]
    tax_rows.sort(key=lambda item: abs(item["difference"]), reverse=True)
    for index, paragraph_index in enumerate((4, 5)):
        if index >= len(tax_rows) or paragraph_index >= len(document.paragraphs):
            continue
        item = tax_rows[index]
        direction = "рост" if item["difference"] >= 0 else "снижение"
        _replace_paragraph(
            document.paragraphs[paragraph_index],
            f"{item['name']}: {direction} на "
            f"{_report_number(abs(item['difference']))} млрд рублей, темп "
            f"{ru_number(item['growth_rate'])}%.",
        )
    leaders = [
        item for item in data["industries"] if (item["growth_rate"] or 0) > 120
    ][:3]
    largest = sorted(data["industries"], key=lambda item: item["share"] or 0, reverse=True)[
        :3
    ]
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Лидеры роста поступлений"):
            text = (
                "Лидеры роста поступлений (темп более 120%): "
                + "; ".join(
                    f"{item['name']} — {ru_number(item['growth_rate'])}%"
                    for item in leaders
                )
                + "."
                if leaders
                else "Лидеры роста поступлений с темпом более 120% не выявлены."
            )
            _replace_paragraph(paragraph, text)
        elif paragraph.text.startswith("Основные отрасли"):
            _replace_paragraph(
                paragraph,
                "Основные отрасли по доле поступлений МСП: "
                + "; ".join(
                    f"{item['name']} — {_report_number(item['share'])}%"
                    for item in largest
                )
                + ".",
            )


def _apply_report_labels(document: Document, data: dict) -> None:
    if data.get("report_mode") != "region_vs_rf":
        return
    territory = data["territory"]
    replacements = {
        0: f"А. Поступления от МСП: {territory}",
        2: (
            "ТАБЛИЦА: СРАВНЕНИЕ СОВОКУПНЫХ ПОСТУПЛЕНИЙ ОТ МСП "
            "(КОНСОЛИДИРОВАННЫЙ БЮДЖЕТ + СТРАХОВЫЕ ВЗНОСЫ): РФ И "
            f"«{territory.upper()}», МЛРД РУБЛЕЙ"
        ),
        16: (
            "ТАБЛИЦА: СРАВНЕНИЕ СОВОКУПНЫХ ПОСТУПЛЕНИЙ ОТ МСП "
            "(КОНСОЛИДИРОВАННЫЙ БЮДЖЕТ + СТРАХОВЫЕ ВЗНОСЫ): РФ И "
            f"«{territory.upper()}», МЛРД РУБЛЕЙ"
        ),
        17: (
            "ТАБЛИЦА: ПОСТУПЛЕНИЯ ОТ МСП В КОНСОЛИДИРОВАННЫЙ "
            f"БЮДЖЕТ РЕГИОНА «{territory.upper()}» ПО ОСНОВНЫМ НАЛОГАМ "
            "(БЕЗ СТРАХОВЫХ ВЗНОСОВ), МЛРД РУБЛЕЙ"
        ),
        18: (
            "ТАБЛИЦА: ПОСТУПЛЕНИЯ ОТ МСП В КОНСОЛИДИРОВАННЫЙ "
            f"БЮДЖЕТ РЕГИОНА «{territory.upper()}» В РАЗРЕЗЕ ОТРАСЛЕЙ "
            "(БЕЗ СТРАХОВЫХ ВЗНОСОВ), МЛРД РУБЛЕЙ"
        ),
    }
    for index, text in replacements.items():
        if index < len(document.paragraphs):
            _replace_paragraph(document.paragraphs[index], text)


def _ensure_appendices_page_break(document: Document) -> None:
    found = False
    for paragraph in document.paragraphs:
        if paragraph.text.strip().casefold() == "приложения":
            paragraph.paragraph_format.page_break_before = True
            found = True
    if not found:
        raise ValueError("В шаблоне не найден заголовок «Приложения»")


def _remove_tags(document: Document) -> None:
    for paragraph in document.paragraphs:
        if TAG_RE.search(paragraph.text):
            _replace_paragraph(paragraph, TAG_RE.sub("", paragraph.text))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = _cell_text(cell)
                if TAG_RE.search(text):
                    _set_cell(cell, TAG_RE.sub("", text).strip())


def extract_report_document(
    source: Path | DocumentType,
    *,
    version: int = 1,
) -> dict:
    """Extract paragraphs and tables in their actual DOCX body order."""
    document = Document(source) if isinstance(source, Path) else source
    blocks: list[dict] = []
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            blocks.append(
                {
                    "id": f"paragraph-{paragraph_index}",
                    "type": "paragraph",
                    "text": paragraph.text,
                    "style": paragraph.style.name
                    if paragraph.style is not None
                    else None,
                    "editable": True,
                }
            )
            paragraph_index += 1
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            blocks.append(
                {
                    "id": f"table-{table_index}",
                    "type": "table",
                    "rows": [
                        [_cell_text(cell) for cell in row.cells]
                        for row in table.rows
                    ],
                    "editable": False,
                }
            )
            table_index += 1
    return {
        "version": version,
        "blocks": blocks,
        "paragraph_count": paragraph_index,
        "table_count": table_index,
    }


def _number_tokens(text: str) -> list[str]:
    return [
        token.replace("\u00a0", " ").replace(" ", "").replace(",", ".")
        for token in NUMBER_RE.findall(text)
    ]


def validate_text_block_edit(original: str, candidate: str) -> None:
    """Text edits may rephrase prose but cannot alter official numbers."""
    if _number_tokens(original) != _number_tokens(candidate):
        raise ValueError(
            "Числовые значения текстового блока нельзя изменять или удалять. "
            "Используйте подтверждённые значения текущего снимка фактов."
        )


def update_report_document(
    docx_path: Path,
    current: dict,
    changed_text_blocks: dict[str, str],
) -> dict:
    known = {
        block["id"]: block
        for block in current.get("blocks", [])
        if block.get("type") == "paragraph"
    }
    for block_id, text in changed_text_blocks.items():
        block = known.get(block_id)
        if block is None or not block.get("editable"):
            raise KeyError(f"Текстовый блок {block_id} не найден")
        validate_text_block_edit(str(block.get("text", "")), text)

    document = Document(docx_path)
    paragraph_index = 0
    for child in document.element.body.iterchildren():
        if not isinstance(child, CT_P):
            continue
        block_id = f"paragraph-{paragraph_index}"
        if block_id in changed_text_blocks:
            _replace_paragraph(
                Paragraph(child, document),
                changed_text_blocks[block_id],
            )
        paragraph_index += 1

    _ensure_appendices_page_break(document)

    temporary_path = docx_path.with_name(
        f".{docx_path.name}.{uuid.uuid4().hex}.tmp"
    )
    try:
        document.save(temporary_path)
        os.replace(temporary_path, docx_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return extract_report_document(
        docx_path,
        version=int(current.get("version", 1)) + 1,
    )


def generate_report(
    template_path: Path,
    output_dir: Path,
    old_path: Path,
    new_path: Path,
    receipts_path: Path,
    codes_path: Path,
    *,
    old_year: int = 2025,
    new_year: int = 2026,
    months_count: int = 6,
) -> tuple[Path, Path, dict]:
    data = build_report_data(
        old_path,
        new_path,
        receipts_path,
        codes_path,
        old_year=old_year,
        new_year=new_year,
        months_count=months_count,
    )
    return generate_report_from_data(
        template_path,
        output_dir,
        data,
        old_year=old_year,
        new_year=new_year,
        months_count=months_count,
    )


def generate_report_from_data(
    template_path: Path,
    output_dir: Path,
    data: dict,
    *,
    old_year: int,
    new_year: int,
    months_count: int,
) -> tuple[Path, Path, dict]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    territory_suffix = ""
    if data.get("report_mode") == "region_vs_rf":
        safe_territory = re.sub(r"[^0-9A-Za-zА-Яа-яЁё]+", "_", data["territory"])
        territory_suffix = f"_{safe_territory.strip('_')}"
    docx_path = output_dir / (
        f"Поступления_МСП{territory_suffix}_{months_count}мес_"
        f"{new_year}_{stamp}.docx"
    )
    json_path = docx_path.with_suffix(".facts.json")
    document = Document(template_path)
    _apply_report_labels(document, data)
    for table in document.tables:
        if _table_contains(table, "{tab1_col_1}"):
            _fill_table1(
                table,
                data["table1"],
                old_year,
                new_year,
                months_count,
                use_data_labels=(data.get("report_mode") == "region_vs_rf"),
            )
        elif _table_contains(table, "{tab2_col_1}"):
            _fill_table2(
                table,
                data["table2"],
                old_year,
                new_year,
                months_count,
                regional=(data.get("report_mode") == "region_vs_rf"),
            )
        elif _table_contains(table, "{tab3_col2}"):
            _fill_table3(table, data, old_year, new_year, months_count)
    _fill_narrative(document, data)
    _remove_tags(document)
    _ensure_appendices_page_break(document)
    document.save(docx_path)
    snapshot = json_ready(data)
    snapshot["report_document"] = extract_report_document(
        document, version=1
    )
    json_path.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return docx_path, json_path, snapshot


def document_has_tags(path: Path) -> bool:
    document = Document(path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        texts.extend(_cell_text(cell) for row in table.rows for cell in row.cells)
    return any(TAG_RE.search(text) for text in texts)
