from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from python_calamine import CalamineWorkbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from .catalog import (
    HeaderMetricResolution,
    METRIC_BY_COLUMN,
    resolve_header_metrics,
)
from .models import (
    ImportBatch,
    MetricDefinition,
    MspFact,
    RowDefinition,
    SchemaMetricMapping,
    SourceFile,
    Territory,
    ValidationResult,
)
from .workbooks import (
    DATE_RE,
    decimal_value,
    normalize_text,
    sha256_file,
)


APPROVED_REGIONAL_COUNTS = {26, 27, 32, 43, 44}
AS_OF_DIR_RE = re.compile(r"^(\d{2})[.](\d{2})[.](\d{4})$")


@dataclass(frozen=True)
class RegionalLayout:
    as_of_date: date
    numeric_count: int
    metric_columns: dict[int, int]
    raw_headers: dict[int, str]
    header_row: int
    territory_column: int
    sheet_names: list[str]
    rows: list[tuple[Any, ...]]


@dataclass(frozen=True)
class RegionalValue:
    territory: str
    territory_type: str
    value: Decimal
    source: dict


@dataclass(frozen=True)
class RegionalMetricResolution:
    resolution: HeaderMetricResolution
    mapping_method: str
    header_conflicts: dict[int, tuple[str, str]]
    comparable_country_values: int = 0
    matching_country_values: int = 0


def _is_ooxml(path: Path) -> bool:
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        return True
    with path.open("rb") as handle:
        return handle.read(4).startswith(b"PK")


def _xlsx_rows(path: Path, sheet_name: str | None = None) -> tuple[
    list[str], str, list[tuple[Any, ...]]
]:
    handle = path.open("rb") if path.suffix.casefold() == ".xls" else None
    try:
        book = load_workbook(
            handle or path, data_only=True, read_only=True
        )
    except Exception as exc:
        if handle is not None:
            handle.close()
        raise ValueError(
            f"Региональная книга физически повреждена: {path.name}"
        ) from exc
    try:
        names = list(book.sheetnames)
        selected = sheet_name or names[0]
        if selected not in names:
            raise ValueError(f"В региональной книге нет листа {selected}")
        rows = [tuple(row) for row in book[selected].iter_rows(values_only=True)]
        return names, selected, rows
    finally:
        book.close()
        if handle is not None:
            handle.close()


def _xls_rows(path: Path, sheet_name: str | None = None) -> tuple[
    list[str], str, list[tuple[Any, ...]]
]:
    try:
        book = xlrd.open_workbook(path, on_demand=True)
    except Exception as xlrd_error:
        try:
            fallback = CalamineWorkbook.from_path(str(path))
            names = list(fallback.sheet_names)
            selected = sheet_name or names[0]
            sheet = fallback.get_sheet_by_name(selected)
            rows = [tuple(row) for row in sheet.to_python()]
            fallback.close()
            return names, selected, rows
        except Exception as calamine_error:
            raise ValueError(
                f"Региональная книга физически повреждена: {path.name}"
            ) from calamine_error
    try:
        names = book.sheet_names()
        selected = sheet_name or names[0]
        if selected not in names:
            raise ValueError(f"В региональной книге нет листа {selected}")
        sheet = book.sheet_by_name(selected)
        rows = [
            tuple(sheet.cell_value(row, column) for column in range(sheet.ncols))
            for row in range(sheet.nrows)
        ]
        return names, selected, rows
    finally:
        book.release_resources()


def _workbook_rows(path: Path, sheet_name: str | None = None):
    if _is_ooxml(path):
        return _xlsx_rows(path, sheet_name)
    return _xls_rows(path, sheet_name)


def _number(value: Any) -> int | None:
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _layout(
    path: Path,
    *,
    sheet_name: str | None = None,
    fallback_as_of: date | None = None,
) -> tuple[RegionalLayout, str]:
    names, selected, rows = _workbook_rows(path, sheet_name)
    return _layout_from_rows(
        names,
        selected,
        rows,
        fallback_as_of=fallback_as_of,
    )


def _layout_from_rows(
    names: list[str],
    selected: str,
    rows: list[tuple[Any, ...]],
    *,
    fallback_as_of: date | None = None,
) -> tuple[RegionalLayout, str]:
    """Validate an already loaded sheet without reopening its workbook."""
    best_row = -1
    best_mapping: dict[int, int] = {}
    for row_index, row in enumerate(rows[:30]):
        mapping = {
            number: column + 1
            for column, value in enumerate(row[:80])
            if (number := _number(value)) is not None
            and 1 <= number <= 100
        }
        if len(mapping) > len(best_mapping) and all(
            item in mapping for item in range(1, min(5, len(mapping)) + 1)
        ):
            best_row, best_mapping = row_index, mapping
    if len(best_mapping) < 4:
        raise ValueError("Не найдена строка номеров граф региональной формы")
    numeric_count = len(best_mapping)
    if numeric_count not in APPROVED_REGIONAL_COUNTS:
        raise ValueError(
            f"Региональная схема содержит неподдерживаемые {numeric_count} граф"
        )

    found_dates: set[date] = set()
    for row in rows[:25]:
        for value in row[:12]:
            if not isinstance(value, str):
                continue
            for match in DATE_RE.finditer(value):
                day, month, year = map(int, match.groups())
                try:
                    found_dates.add(date(year, month, day))
                except ValueError:
                    pass
    if len(found_dates) == 1:
        as_of = next(iter(found_dates))
    elif fallback_as_of is not None:
        as_of = fallback_as_of
    else:
        raise ValueError("Не удалось однозначно определить дату региональной книги")

    first_metric = best_mapping[min(best_mapping)]
    raw_headers: dict[int, str] = {}
    for source_column, physical_column in sorted(best_mapping.items()):
        parts: list[str] = []
        for header_index in range(max(0, best_row - 6), best_row):
            row = rows[header_index]
            value = (
                row[physical_column - 1]
                if len(row) >= physical_column
                else None
            )
            normalized = normalize_text(value)
            if normalized and normalized not in parts:
                parts.append(normalized)
        raw_headers[source_column] = " → ".join(parts)
    resolution = resolve_header_metrics(raw_headers)
    if not resolution.is_complete:
        raise ValueError(
            "Семантическое сопоставление региональных заголовков "
            f"не прошло: не распознаны {sorted(resolution.unmapped)}, "
            f"дубли {resolution.duplicate_codes}"
        )
    # Current layouts use A for territory; keeping this derived fallback makes
    # the adapter tolerate an inserted service column.
    territory_column = max(1, first_metric - 3)
    return (
        RegionalLayout(
            as_of_date=as_of,
            numeric_count=numeric_count,
            metric_columns=best_mapping,
            raw_headers=raw_headers,
            header_row=best_row + 1,
            territory_column=territory_column,
            sheet_names=names,
            rows=rows,
        ),
        selected,
    )


def directory_as_of(path: Path) -> date | None:
    match = AS_OF_DIR_RE.match(path.parent.name)
    if not match:
        return None
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def inspect_regional_workbook(path: Path) -> RegionalLayout:
    layout, _ = _layout(path, fallback_as_of=directory_as_of(path))
    return layout


def resolve_regional_metrics(
    session: Session,
    *,
    source: SourceFile,
    layout: RegionalLayout,
    section: str,
    row_code: int,
) -> RegionalMetricResolution:
    """Resolve regional metrics and verify the layout against the RF control row.

    A released 2026 regional workbook has headers for a newer 43-column
    template, but its data rows are copied column-for-column from the
    corresponding national 44-column form.  A header-only resolver therefore
    shifts four insurance metrics.  We accept a corresponding national schema
    only when the country row proves the physical alignment with at least four
    comparable values and at least 90% exact matches.
    """

    local = resolve_header_metrics(layout.raw_headers)
    fallback = RegionalMetricResolution(
        resolution=local,
        mapping_method="normalized_header",
        header_conflicts={},
    )
    national_source = session.scalar(
        select(SourceFile)
        .where(
            SourceFile.workbook_kind == "msp_summary",
            SourceFile.status.in_(["valid", "valid_with_warnings"]),
            SourceFile.period_end == source.period_end,
        )
        .order_by(SourceFile.id.desc())
    )
    if national_source is None or national_source.schema_version_id is None:
        return fallback

    national_mapping = {
        report_column: semantic_code
        for report_column, semantic_code in session.execute(
            select(
                SchemaMetricMapping.report_column_number,
                MetricDefinition.semantic_code,
            )
            .join(
                MetricDefinition,
                MetricDefinition.id == SchemaMetricMapping.metric_id,
            )
            .where(
                SchemaMetricMapping.schema_version_id
                == national_source.schema_version_id,
                SchemaMetricMapping.approved.is_(True),
            )
        )
    }
    if not all(
        column in national_mapping for column in layout.metric_columns
    ):
        return fallback

    country_row = next(
        (
            row
            for row in layout.rows[layout.header_row :]
            if len(row) >= layout.territory_column
            and territory_type(
                str(row[layout.territory_column - 1] or "").strip()
            )
            == "country"
        ),
        None,
    )
    if country_row is None:
        return fallback

    national_facts = session.execute(
        select(MspFact, MetricDefinition.semantic_code)
        .join(
            MetricDefinition,
            MetricDefinition.id == MspFact.metric_id,
        )
        .join(
            RowDefinition,
            RowDefinition.id == MspFact.row_definition_id,
        )
        .where(
            MspFact.source_file_id == national_source.id,
            MspFact.section == section,
            RowDefinition.base_row_code == row_code,
        )
    ).all()
    national_values_by_source: dict[int, Decimal] = {}
    for fact, _semantic_code in national_facts:
        source_column = (fact.lineage_json or {}).get(
            "source_report_column"
        )
        if isinstance(source_column, int):
            national_values_by_source[source_column] = fact.value

    comparable = 0
    matches = 0
    for source_column, physical_column in layout.metric_columns.items():
        if (
            source_column not in national_values_by_source
            or len(country_row) < physical_column
        ):
            continue
        regional_value = decimal_value(country_row[physical_column - 1])
        if regional_value is None:
            continue
        comparable += 1
        if regional_value == national_values_by_source[source_column]:
            matches += 1
    if comparable < 4 or matches / comparable < 0.9:
        return fallback

    aligned = HeaderMetricResolution(
        source_to_code={
            column: national_mapping[column]
            for column in sorted(layout.metric_columns)
        },
        unmapped={},
        duplicate_codes={},
    )
    conflicts = {
        column: (
            local.source_to_code[column],
            aligned.source_to_code[column],
        )
        for column in aligned.source_to_code
        if local.source_to_code.get(column)
        and local.source_to_code[column] != aligned.source_to_code[column]
    }
    return RegionalMetricResolution(
        resolution=aligned,
        mapping_method="national_schema_country_crosscheck",
        header_conflicts=conflicts,
        comparable_country_values=comparable,
        matching_country_values=matches,
    )


def territory_type(name: str) -> str | None:
    normalized = normalize_text(name)
    if not normalized or normalized == "в том числе:":
        return None
    if normalized == "российская федерация":
        return "country"
    if "федеральный округ" in normalized:
        return "federal_district"
    if normalized == "юго-западные субъекты":
        return "territory_group"
    return "region"


def territory_key(name: str) -> str:
    normalized = normalize_text(name)
    # Regional workbooks use both "г. Москва" and "город Москва" (the same
    # also happens for other federal cities).  Treat the presentation prefix
    # as an alias so one user mention never becomes two required territories.
    normalized = re.sub(r"^(?:город\s+|г[.]?\s*)", "", normalized)
    # The source inventory contains long official names while some regional
    # books use their statutory short forms.  These are presentation aliases,
    # not different territories.  Canonicalise only the six variants that are
    # present in the verified inventory; broader fuzzy matching here could
    # silently merge genuinely different regions.
    normalized = re.sub(r"\s*\([^()]+\)\s*$", "", normalized)
    normalized = re.sub(r"\s*-\s*кузбасс\s*$", "", normalized)
    normalized = re.sub(r"\bавтономный\s+округ\b", "ао", normalized)
    return re.sub(r"[^0-9a-zа-я]+", "", normalized)


def register_regional_file(
    session: Session,
    batch: ImportBatch,
    path: Path,
    kind: str,
) -> SourceFile:
    layout = inspect_regional_workbook(path)
    digest = sha256_file(path)
    source = session.scalar(
        select(SourceFile).where(SourceFile.sha256 == digest)
    )
    if source is None:
        source = SourceFile(
            batch_id=batch.id,
            absolute_path=str(path.resolve()),
            sha256=digest,
            size_bytes=path.stat().st_size,
            modified_at=datetime.fromtimestamp(path.stat().st_mtime),
            file_format=path.suffix.casefold().lstrip("."),
            workbook_kind=kind,
        )
        session.add(source)
        session.flush()
    period_end = layout.as_of_date - timedelta(days=1)
    source.workbook_kind = kind
    source.as_of_date = layout.as_of_date
    source.period_start = date(period_end.year, 1, 1)
    source.period_end = period_end
    source.report_year = period_end.year
    source.months_count = period_end.month
    source.series_kind = "cumulative_ytd"
    source.status = "valid"
    source.warnings_json = [
        "Региональные факты читаются проверенным потоковым адаптером "
        "по выбранной строке формы."
    ]
    session.execute(
        delete(ValidationResult).where(
            ValidationResult.source_file_id == source.id,
            ValidationResult.check_code == "unsupported_workbook_layout",
        )
    )

    if kind == "regional_section_i":
        for row in layout.rows[layout.header_row :]:
            raw_name = (
                row[layout.territory_column - 1]
                if len(row) >= layout.territory_column
                else None
            )
            name = str(raw_name or "").strip()
            kind_name = territory_type(name)
            if kind_name is None:
                continue
            normalized = normalize_text(name)
            territory = session.scalar(
                select(Territory).where(
                    Territory.normalized_name == normalized
                )
            )
            if territory is None:
                session.add(
                    Territory(
                        name_ru=name,
                        normalized_name=normalized,
                        territory_type=kind_name,
                    )
                )
    session.flush()
    return source


def _source_for(
    session: Session,
    period: date,
    section: str,
    row_code: int,
) -> SourceFile:
    kind = (
        "regional_section_i" if section == "I" else "regional_section_ii"
    )
    sources = session.scalars(
        select(SourceFile)
        .where(SourceFile.workbook_kind == kind)
        .order_by(SourceFile.id.desc())
    ).all()
    sheet_code = row_code + (1000 if section == "II" else 0)
    for source in sources:
        source_period = source.period_end
        if source_period is None:
            inferred = directory_as_of(Path(source.absolute_path))
            source_period = inferred - timedelta(days=1) if inferred else None
        if source_period != period:
            continue
        try:
            layout = inspect_regional_workbook(Path(source.absolute_path))
        except ValueError:
            continue
        if str(sheet_code) in layout.sheet_names:
            return source
    raise ValueError(
        f"Нет региональной книги Раздела {section} за {period}, "
        f"содержащей отраслевой лист {sheet_code}"
    )


def read_regional_values(
    session: Session,
    *,
    period: date,
    section: str,
    row_code: int,
    report_column: int,
) -> list[RegionalValue]:
    source = _source_for(session, period, section, row_code)
    path = Path(source.absolute_path)
    sheet_code = row_code + (1000 if section == "II" else 0)
    layout, selected = _layout(
        path,
        sheet_name=str(sheet_code),
        fallback_as_of=source.as_of_date or directory_as_of(path),
    )
    metric_seed = METRIC_BY_COLUMN.get(report_column)
    if metric_seed is None:
        raise ValueError(
            f"В каталоге нет семантики для показателя {report_column}"
        )
    metric_resolution = resolve_regional_metrics(
        session,
        source=source,
        layout=layout,
        section=section,
        row_code=row_code,
    )
    resolution = metric_resolution.resolution
    source_column = next(
        (
            column
            for column, code in resolution.source_to_code.items()
            if code == metric_seed.code
        ),
        None,
    )
    physical_column = (
        layout.metric_columns.get(source_column)
        if source_column is not None
        else None
    )
    if physical_column is None:
        raise ValueError(
            "В региональной схеме нет показателя "
            f"«{metric_seed.name}» ({metric_seed.code})"
        )
    values: list[RegionalValue] = []
    for row_number, row in enumerate(
        layout.rows[layout.header_row :],
        start=layout.header_row + 1,
    ):
        if len(row) < max(layout.territory_column, physical_column):
            continue
        name = str(row[layout.territory_column - 1] or "").strip()
        kind_name = territory_type(name)
        value = decimal_value(row[physical_column - 1])
        if kind_name is None or value is None:
            continue
        values.append(
            RegionalValue(
                territory=name,
                territory_type=kind_name,
                value=value,
                source={
                    "file": path.name,
                    "absolute_path": str(path),
                    "sheet": selected,
                    "cell": f"{get_column_letter(physical_column)}{row_number}",
                    "row_code": sheet_code,
                    "report_column": source_column,
                    "canonical_report_column": report_column,
                    "semantic_code": metric_seed.code,
                    "source_header": layout.raw_headers.get(
                        source_column,
                        "",
                    ),
                    "mapping_method": metric_resolution.mapping_method,
                    "header_conflict": metric_resolution.header_conflicts.get(
                        source_column
                    ),
                    "country_crosscheck": {
                        "comparable": (
                            metric_resolution.comparable_country_values
                        ),
                        "matching": metric_resolution.matching_country_values,
                    },
                    "raw_value": str(value),
                    "read_only": True,
                },
            )
        )
    return values


def read_regional_industry_values(
    session: Session,
    *,
    period: date,
    section: str,
    row_codes: list[int],
    report_column: int,
    territory: str,
) -> dict[int, RegionalValue]:
    """Read one territory across many industry sheets in one workbook pass.

    Opening a 378-sheet regional XLSX once per row takes several minutes.
    This adapter opens the verified source once, validates every selected
    sheet against the workbook's approved column layout, and returns only the
    requested territory.
    """

    kind = (
        "regional_section_i" if section == "I" else "regional_section_ii"
    )
    source = session.scalar(
        select(SourceFile)
        .where(
            SourceFile.workbook_kind == kind,
            SourceFile.period_end == period,
            SourceFile.status.in_(["valid", "valid_with_warnings"]),
        )
        .order_by(SourceFile.id.desc())
    )
    if source is None:
        raise ValueError(
            f"Нет региональной книги Раздела {section} за {period}"
        )
    metric_seed = METRIC_BY_COLUMN.get(report_column)
    if metric_seed is None:
        raise ValueError(
            f"В каталоге нет семантики для показателя {report_column}"
        )
    path = Path(source.absolute_path)
    try:
        book = CalamineWorkbook.from_path(str(path))
    except Exception as exc:
        raise ValueError(
            f"Региональная книга физически повреждена: {path.name}"
        ) from exc
    try:
        names = list(book.sheet_names)
        requested = [
            (row_code, str(row_code + (1000 if section == "II" else 0)))
            for row_code in row_codes
        ]
        available = [(code, sheet) for code, sheet in requested if sheet in names]
        if not available:
            return {}

        first_code, first_sheet = available[0]
        first_rows = [
            tuple(row)
            for row in book.get_sheet_by_name(first_sheet).to_python()
        ]
        base_layout, _ = _layout_from_rows(
            names,
            first_sheet,
            first_rows,
            fallback_as_of=(
                source.as_of_date or directory_as_of(path)
            ),
        )
        metric_resolution = resolve_regional_metrics(
            session,
            source=source,
            layout=base_layout,
            section=section,
            row_code=first_code,
        )
        source_column = next(
            (
                column
                for column, code in (
                    metric_resolution.resolution.source_to_code.items()
                )
                if code == metric_seed.code
            ),
            None,
        )
        physical_column = (
            base_layout.metric_columns.get(source_column)
            if source_column is not None
            else None
        )
        if physical_column is None:
            raise ValueError(
                "В региональной схеме нет показателя "
                f"«{metric_seed.name}» ({metric_seed.code})"
            )

        wanted_territory = territory_key(territory)
        results: dict[int, RegionalValue] = {}
        for row_code, sheet_name in available:
            rows = (
                first_rows
                if sheet_name == first_sheet
                else [
                    tuple(row)
                    for row in book.get_sheet_by_name(sheet_name).to_python()
                ]
            )
            layout, selected = _layout_from_rows(
                names,
                sheet_name,
                rows,
                fallback_as_of=(
                    source.as_of_date or directory_as_of(path)
                ),
            )
            if layout.metric_columns != base_layout.metric_columns:
                raise ValueError(
                    "Региональные отраслевые листы имеют разные схемы "
                    f"граф: {first_sheet} и {sheet_name}"
                )
            for row_number, row in enumerate(
                layout.rows[layout.header_row :],
                start=layout.header_row + 1,
            ):
                if len(row) < max(
                    layout.territory_column, physical_column
                ):
                    continue
                name = str(
                    row[layout.territory_column - 1] or ""
                ).strip()
                if territory_key(name) != wanted_territory:
                    continue
                value = decimal_value(row[physical_column - 1])
                if value is None:
                    break
                sheet_code = row_code + (1000 if section == "II" else 0)
                results[row_code] = RegionalValue(
                    territory=name,
                    territory_type=territory_type(name) or "region",
                    value=value,
                    source={
                        "file": path.name,
                        "absolute_path": str(path),
                        "sheet": selected,
                        "cell": (
                            f"{get_column_letter(physical_column)}"
                            f"{row_number}"
                        ),
                        "row_code": sheet_code,
                        "report_column": source_column,
                        "canonical_report_column": report_column,
                        "semantic_code": metric_seed.code,
                        "source_header": layout.raw_headers.get(
                            source_column, ""
                        ),
                        "mapping_method": (
                            metric_resolution.mapping_method
                        ),
                        "header_conflict": (
                            metric_resolution.header_conflicts.get(
                                source_column
                            )
                        ),
                        "country_crosscheck": {
                            "comparable": (
                                metric_resolution.comparable_country_values
                            ),
                            "matching": (
                                metric_resolution.matching_country_values
                            ),
                        },
                        "raw_value": str(value),
                        "read_only": True,
                    },
                )
                break
        return results
    finally:
        book.close()


def read_regional_report_values(
    session: Session,
    *,
    period: date,
    section: str,
    row_codes: list[int],
    report_columns: list[int],
    territory: str,
    report_columns_by_row: dict[int, list[int]] | None = None,
) -> dict[tuple[int, int], RegionalValue]:
    """Read all standard-report cells from one regional workbook pass.

    The returned key is ``(base_row_code, canonical_report_column)``.  Every
    requested sheet, metric and territory must be present: a standard report
    must never replace a missing period or cell with a nearby value.
    """
    if section not in {"I", "II"}:
        raise ValueError("Региональный reader поддерживает только разделы I/II")
    unique_rows = list(dict.fromkeys(row_codes))
    unique_columns = list(dict.fromkeys(report_columns))
    if not unique_rows or not unique_columns:
        return {}

    kind = (
        "regional_section_i" if section == "I" else "regional_section_ii"
    )
    source = session.scalar(
        select(SourceFile)
        .where(
            SourceFile.workbook_kind == kind,
            SourceFile.period_end == period,
            SourceFile.status.in_(["valid", "valid_with_warnings"]),
        )
        .order_by(SourceFile.id.desc())
    )
    if source is None:
        raise ValueError(
            f"Нет региональной книги Раздела {section} за {period.isoformat()}"
        )

    path = Path(source.absolute_path)
    try:
        book = CalamineWorkbook.from_path(str(path))
    except Exception as exc:
        raise ValueError(
            f"Региональная книга физически повреждена: {path.name}"
        ) from exc
    try:
        names = list(book.sheet_names)
        requested_sheets = {
            row_code: str(row_code + (1000 if section == "II" else 0))
            for row_code in unique_rows
        }
        missing_sheets = [
            sheet for sheet in requested_sheets.values() if sheet not in names
        ]
        if missing_sheets:
            raise ValueError(
                f"В региональной книге Раздела {section} за "
                f"{period.isoformat()} отсутствуют листы: "
                + ", ".join(missing_sheets)
            )

        first_row_code = unique_rows[0]
        first_sheet = requested_sheets[first_row_code]
        first_rows = [
            tuple(row)
            for row in book.get_sheet_by_name(first_sheet).to_python()
        ]
        base_layout, _ = _layout_from_rows(
            names,
            first_sheet,
            first_rows,
            fallback_as_of=source.as_of_date or directory_as_of(path),
        )
        metric_resolution = resolve_regional_metrics(
            session,
            source=source,
            layout=base_layout,
            section=section,
            row_code=first_row_code,
        )

        resolved_columns: dict[int, tuple[int, int]] = {}
        for canonical_column in unique_columns:
            metric = METRIC_BY_COLUMN.get(canonical_column)
            if metric is None:
                raise ValueError(
                    "В каталоге нет семантики для показателя "
                    f"{canonical_column}"
                )
            source_column = next(
                (
                    column
                    for column, code in (
                        metric_resolution.resolution.source_to_code.items()
                    )
                    if code == metric.code
                ),
                None,
            )
            physical_column = (
                base_layout.metric_columns.get(source_column)
                if source_column is not None
                else None
            )
            if source_column is None or physical_column is None:
                raise ValueError(
                    f"В региональной схеме нет показателя «{metric.name}» "
                    f"({metric.code})"
                )
            resolved_columns[canonical_column] = (
                source_column,
                physical_column,
            )

        wanted_territory = territory_key(territory)
        result: dict[tuple[int, int], RegionalValue] = {}
        for row_code, sheet_name in requested_sheets.items():
            rows = (
                first_rows
                if sheet_name == first_sheet
                else [
                    tuple(row)
                    for row in book.get_sheet_by_name(sheet_name).to_python()
                ]
            )
            layout, selected = _layout_from_rows(
                names,
                sheet_name,
                rows,
                fallback_as_of=source.as_of_date or directory_as_of(path),
            )
            if layout.metric_columns != base_layout.metric_columns:
                raise ValueError(
                    "Региональные листы имеют разные схемы граф: "
                    f"{first_sheet} и {sheet_name}"
                )
            matching_row: tuple[int, tuple[Any, ...], str] | None = None
            for row_number, row in enumerate(
                layout.rows[layout.header_row :],
                start=layout.header_row + 1,
            ):
                if len(row) < layout.territory_column:
                    continue
                source_territory = str(
                    row[layout.territory_column - 1] or ""
                ).strip()
                if territory_key(source_territory) == wanted_territory:
                    matching_row = (row_number, row, source_territory)
                    break
            if matching_row is None:
                raise ValueError(
                    f"Регион «{territory}» не найден в {path.name}!{sheet_name}"
                )
            row_number, row, source_territory = matching_row
            wanted_columns = (
                report_columns_by_row.get(row_code, unique_columns)
                if report_columns_by_row is not None
                else unique_columns
            )
            for canonical_column in wanted_columns:
                source_column, physical_column = resolved_columns[
                    canonical_column
                ]
                value = (
                    decimal_value(row[physical_column - 1])
                    if len(row) >= physical_column
                    else None
                )
                metric = METRIC_BY_COLUMN[canonical_column]
                if value is None:
                    raise ValueError(
                        f"Нет значения «{metric.name}» для региона "
                        f"«{territory}» в {path.name}!{sheet_name}"
                    )
                actual_row_code = row_code + (1000 if section == "II" else 0)
                result[(row_code, canonical_column)] = RegionalValue(
                    territory=source_territory,
                    territory_type=territory_type(source_territory) or "region",
                    value=value,
                    source={
                        "file": path.name,
                        "absolute_path": str(path.resolve()),
                        "sheet": selected,
                        "cell": (
                            f"{get_column_letter(physical_column)}{row_number}"
                        ),
                        "row_code": actual_row_code,
                        "report_column": source_column,
                        "source_report_column": source_column,
                        "canonical_report_column": canonical_column,
                        "report_column_name": metric.name,
                        "semantic_code": metric.code,
                        "source_header": layout.raw_headers.get(
                            source_column, ""
                        ),
                        "mapping_method": metric_resolution.mapping_method,
                        "header_conflict": (
                            metric_resolution.header_conflicts.get(source_column)
                        ),
                        "country_crosscheck": {
                            "comparable": (
                                metric_resolution.comparable_country_values
                            ),
                            "matching": (
                                metric_resolution.matching_country_values
                            ),
                        },
                        "section": section,
                        "territory": source_territory,
                        "period_end": period.isoformat(),
                        "raw_value": str(value),
                        "raw_unit": (
                            "ед." if canonical_column == 1 else "тыс. руб."
                        ),
                        "conversion": (
                            None
                            if canonical_column == 1
                            else "тыс. руб. ÷ 1 000 000 = млрд руб."
                        ),
                        "read_only": True,
                    },
                )
        return result
    finally:
        book.close()
