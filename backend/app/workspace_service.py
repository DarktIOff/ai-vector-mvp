from __future__ import annotations

import os
import re
import textwrap
import uuid
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .config import Settings
from .clarifications import resolve_clarification_answer
from .models import (
    AnalysisProject,
    AnalysisProjectBlock,
    Conversation,
    ConversationTurn,
    Scenario,
)
from .planner import PlannerResult, plan_question, validate_plan
from .query_service import available_periods, build_query_plan, execute_plan
from .reporting import validate_text_block_edit
from .schemas import QueryPlan
from .workbooks import normalize_text


PROJECT_PRESENTATION_VERSION = 5

_DOCX_NAVY = "17324D"
_DOCX_ACCENT = "2563EB"
_DOCX_ACCENT_SOFT = "EAF2FF"
_DOCX_TEXT = "1E293B"
_DOCX_MUTED = "64748B"
_DOCX_BORDER = "D7E0E8"
_DOCX_ROW_ALT = "F6F9FC"
_DOCX_SUCCESS = "2E7D5A"
_DOCX_SUCCESS_SOFT = "ECF8F2"
_DOCX_WARNING = "C47A08"
_DOCX_WARNING_SOFT = "FFF7E6"


def _now() -> datetime:
    return datetime.utcnow()


def _new_id() -> str:
    return str(uuid.uuid4())


def _title_from_question(question: str) -> str:
    title = re.sub(r"\s+", " ", question).strip()
    return title if len(title) <= 64 else title[:61].rstrip() + "…"


def conversation_summary(conversation: Conversation) -> dict:
    return {
        "id": conversation.id,
        "title": conversation.title,
        "status": conversation.status,
        "pinned": conversation.pinned,
        "archived": conversation.archived,
        "context": _context_label(conversation.state_json or {}),
        "created_at": conversation.created_at.isoformat(),
        "updated_at": conversation.updated_at.isoformat(),
        "project_id": conversation.project.id if conversation.project else None,
        "turns_count": len(conversation.turns),
    }


def _context_label(state: dict) -> str:
    metrics = state.get("metric_codes") or []
    period = state.get("period")
    territory = state.get("territory")
    parts = []
    if metrics:
        parts.append(
            "Поступления МСП"
            if len(metrics) > 1
            else str(metrics[0]).replace("receipts_", "").replace("_", " ")
        )
    if territory:
        parts.append(str(territory))
    if period:
        parts.append(str(period))
    return " · ".join(parts) or "Новый аналитический диалог"


def turn_record(turn: ConversationTurn) -> dict:
    result = dict(turn.result_json or {}) or None
    if result and result.get("answer"):
        result["answer"] = format_report_text_numbers(result["answer"])
    return {
        "id": turn.id,
        "sequence": turn.sequence,
        "question": turn.content,
        "status": turn.status,
        "turn_type": turn.turn_type,
        "plan": turn.plan_json or None,
        "result": result,
        "clarification": turn.clarification_json or None,
        "planner": turn.planner_json or {},
        "created_at": turn.created_at.isoformat(),
    }


def conversation_detail(conversation: Conversation) -> dict:
    return {
        **conversation_summary(conversation),
        "state": conversation.state_json or {},
        "turns": [turn_record(turn) for turn in conversation.turns],
        "project": project_to_dict(conversation.project)
        if conversation.project
        else None,
    }


def _analysis_entities(turn: ConversationTurn) -> list[dict]:
    """Return source-backed analytical entity identifiers without numbers."""
    result = turn.result_json or {}
    facts_by_id = {
        str(item.get("fact_id")): item
        for item in result.get("facts", [])
        if isinstance(item, dict) and item.get("fact_id") is not None
    }
    entities: dict[int, str] = {}
    primary_sections = {
        "ranking_value",
        "ranking_growth",
        "industry_growth_best",
        "industry_growth_worst",
    }
    for section in result.get("analysis_sections", []):
        if (
            not isinstance(section, dict)
            or section.get("kind") not in primary_sections
        ):
            continue
        for row in section.get("rows", []):
            if not isinstance(row, dict):
                continue
            row_code = row.get("row_code")
            if row_code is None:
                for fact_id in row.get("source_fact_ids", []):
                    fact = facts_by_id.get(str(fact_id))
                    if isinstance(fact, dict) and fact.get("row_code") is not None:
                        row_code = fact["row_code"]
                        break
            if (
                isinstance(row_code, int)
                and row_code != 1010
                and row.get("entity")
            ):
                entities.setdefault(row_code, str(row["entity"]))
    return [
        {"name": name, "row_code": row_code}
        for row_code, name in entities.items()
    ]


def _planner_context(conversation: Conversation) -> list[dict]:
    result = []
    for turn in conversation.turns[-12:]:
        if turn.status == "needs_clarification":
            clarification = turn.clarification_json or {}
            result.append(
                {
                    "question": turn.content,
                    "pending_question": turn.content,
                    "clarification_question": clarification.get("question"),
                    "clarification_options": clarification.get("options", []),
                }
            )
            continue
        if turn.status != "ready" or not turn.plan_json:
            continue
        stored_plan = turn.plan_json
        if (
            isinstance(stored_plan, dict)
            and "operations" in stored_plan
            and isinstance(turn.result_json, dict)
        ):
            stored_plan = turn.result_json.get("query_plan") or {}
        result.append(
            {
                "question": turn.content,
                "plan": stored_plan,
                "result_summary": {
                    "expected_analyses": list(
                        (
                            (turn.result_json or {}).get("research_plan")
                            or {}
                        ).get("expected_analyses", [])
                    ),
                    "analysis_kinds": [
                        item.get("kind")
                        for item in (turn.result_json or {}).get(
                            "analysis_sections", []
                        )
                    ],
                    "analysis_entities": _analysis_entities(turn),
                },
            }
        )
    return result


def _state_from_plan(plan: QueryPlan) -> dict:
    return {
        "metric_codes": list(plan.metric_codes),
        "enterprise_groups": list(plan.enterprise_groups),
        "territories": list(plan.territories),
        "territory": plan.territories[0] if plan.territories else None,
        "row_codes": list(plan.row_codes),
        "okved_codes": list(plan.okved_codes),
        "period": plan.periods[-1].isoformat() if plan.periods else None,
        "period_mode": plan.period_modes[-1] if plan.period_modes else None,
        "intent": plan.intent,
        "analysis_dimension": plan.analysis_dimension,
        "comparison_basis": plan.comparison_basis,
    }


def _option_id(label: str, index: int) -> str:
    normalized = re.sub(r"[^a-zа-я0-9]+", "-", normalize_text(label)).strip("-")
    return (normalized[:72] or f"option-{index}") + f"-{index}"


def _structured_options(planned: PlannerResult) -> list[dict]:
    question = normalize_text(planned.clarification_question or "")
    if "показател" in question or "поступлен" in question:
        return [
            {
                "id": "budget_receipts",
                "label": "Налоговые поступления",
                "plan_patch": {
                    "metric_codes": ["receipts_budget_total"],
                },
            },
            {
                "id": "insurance_receipts",
                "label": "Страховые взносы",
                "plan_patch": {
                    "metric_codes": ["receipts_insurance_total"],
                },
            },
            {
                "id": "all_msp_receipts",
                "label": "Поступления от МСП, всего",
                "plan_patch": {
                    "metric_codes": [
                        "receipts_budget_total",
                        "receipts_insurance_total",
                    ],
                },
            },
        ]
    result = []
    for index, label in enumerate(planned.clarification_options or [], start=1):
        patch: dict = {}
        normalized = normalize_text(label)
        if "вся россия" in normalized:
            patch["territories"] = ["Российская Федерация"]
        result.append(
            {
                "id": _option_id(label, index),
                "label": label,
                "plan_patch": patch,
            }
        )
    return result


def _response_for_turn(
    conversation: Conversation, turn: ConversationTurn
) -> dict:
    if turn.status == "ready":
        return {
            **(turn.result_json or {}),
            "turn_id": turn.id,
            "sequence": turn.sequence,
            "conversation_id": conversation.id,
            "question": turn.content,
            "status": "ready",
        }
    clarification = turn.clarification_json or {}
    return {
        "turn_id": turn.id,
        "sequence": turn.sequence,
        "query_id": turn.id,
        "conversation_id": conversation.id,
        "question": turn.content,
        "status": "needs_clarification",
        "query_plan": None,
        "resolved_metrics": [],
        "facts": [],
        "calculations": [],
        "series": [],
        "analysis_sections": [],
        "visualization": None,
        "warnings": (turn.planner_json or {}).get("warnings", []),
        "answer": clarification.get("question", "Уточните запрос."),
        "answer_mode": "clarification",
        "numbers_verified": True,
        "clarification_question": clarification.get("question"),
        "clarification_options": clarification.get("options", []),
    }


def _store_ready_result(
    session: Session,
    conversation: Conversation,
    turn: ConversationTurn,
    plan: QueryPlan,
    result: dict,
    planner_mode: str,
    planner_warnings: list[str],
    tool_journal: list[dict] | None,
) -> dict:
    result = dict(result)
    result.update(
        {
            "query_id": turn.id,
            "turn_id": turn.id,
            "conversation_id": conversation.id,
            "question": turn.content,
            "planner_mode": planner_mode,
            "warnings": [
                *planner_warnings,
                *result.get("warnings", []),
            ],
            "status": "ready",
        }
    )
    turn.status = "ready"
    turn.turn_type = "analysis"
    turn.plan_json = plan.model_dump(mode="json")
    turn.result_json = result
    turn.clarification_json = {}
    turn.planner_json = {
        "mode": planner_mode,
        "warnings": planner_warnings,
        "tool_journal": tool_journal or [],
        "validated": True,
    }
    conversation.state_json = _state_from_plan(plan)
    conversation.updated_at = _now()
    if conversation.title == "Новый диалог":
        conversation.title = _title_from_question(turn.content)
    session.commit()
    return _response_for_turn(conversation, turn)


async def create_conversation_turn(
    session: Session,
    settings: Settings,
    conversation: Conversation,
    question: str,
) -> dict:
    sequence = (
        session.scalar(
            select(func.max(ConversationTurn.sequence)).where(
                ConversationTurn.conversation_id == conversation.id
            )
        )
        or 0
    ) + 1
    turn = ConversationTurn(
        id=_new_id(),
        conversation_id=conversation.id,
        sequence=sequence,
        role="user",
        turn_type="question",
        status="processing",
        content=question.strip(),
    )
    session.add(turn)
    conversation.updated_at = _now()
    session.commit()
    session.expire(conversation, ["turns"])
    try:
        planned = await plan_question(
            session,
            settings,
            turn.content,
            available_periods(session),
            _planner_context(conversation),
        )
        if planned.status == "needs_clarification" or planned.plan is None:
            options = _structured_options(planned)
            turn.status = "needs_clarification"
            turn.turn_type = "clarification"
            turn.clarification_json = {
                "question": planned.clarification_question
                or "Уточните параметры запроса.",
                "options": options,
            }
            turn.planner_json = {
                "mode": planned.mode,
                "warnings": planned.warnings,
                "tool_journal": planned.tool_journal or [],
                "validated": False,
            }
            session.commit()
            return _response_for_turn(conversation, turn)
        result = execute_plan(session, planned.plan)
        return _store_ready_result(
            session,
            conversation,
            turn,
            planned.plan,
            result,
            planned.mode,
            planned.warnings,
            planned.tool_journal,
        )
    except Exception:
        turn.status = "failed"
        turn.turn_type = "error"
        session.commit()
        raise


def _plan_from_option(
    session: Session,
    turn: ConversationTurn,
    option: dict,
) -> QueryPlan:
    patch = option.get("plan_patch") or {}
    label = str(option.get("label") or "")
    periods = available_periods(session)
    clarification = turn.clarification_json or {}
    clarified_question = resolve_clarification_answer(
        turn.content,
        label,
        clarification_question=str(clarification.get("question") or ""),
        clarification_options=clarification.get("options"),
    )
    try:
        plan = build_query_plan(clarified_question, periods)
    except ValueError:
        plan = QueryPlan(
            metric_codes=["receipts_budget_total"],
            periods=[max(periods)],
        )
    raw = plan.model_dump(mode="python")
    raw.update(patch)
    return validate_plan(session, QueryPlan.model_validate(raw), periods)


async def answer_clarification(
    session: Session,
    settings: Settings,
    conversation: Conversation,
    turn: ConversationTurn,
    *,
    option_id: str | None,
    text: str | None,
) -> dict:
    if turn.status != "needs_clarification":
        raise ValueError("Этот ход не ожидает уточнения")
    clarification = turn.clarification_json or {}
    options = clarification.get("options") or []
    selected = next(
        (item for item in options if item.get("id") == option_id),
        None,
    )
    if selected is not None:
        plan = _plan_from_option(session, turn, selected)
        result = execute_plan(session, plan)
        return _store_ready_result(
            session,
            conversation,
            turn,
            plan,
            result,
            "clarification_patch",
            [],
            [
                {
                    "action": "apply_clarification_option",
                    "option_id": option_id,
                    "status": "valid",
                }
            ],
        )
    if not text or not text.strip():
        raise ValueError("Выберите вариант или введите уточнение")
    planned = await plan_question(
        session,
        settings,
        text.strip(),
        available_periods(session),
        _planner_context(conversation),
    )
    if planned.status == "needs_clarification" or planned.plan is None:
        turn.clarification_json = {
            "question": planned.clarification_question
            or "Нужно ещё одно конкретное уточнение.",
            "options": _structured_options(planned),
        }
        turn.planner_json = {
            "mode": planned.mode,
            "warnings": planned.warnings,
            "tool_journal": planned.tool_journal or [],
            "validated": False,
        }
        session.commit()
        return _response_for_turn(conversation, turn)
    result = execute_plan(session, planned.plan)
    return _store_ready_result(
        session,
        conversation,
        turn,
        planned.plan,
        result,
        planned.mode,
        planned.warnings,
        planned.tool_journal,
    )


def ensure_project(session: Session, conversation: Conversation) -> AnalysisProject:
    if conversation.project is not None:
        return conversation.project
    project = AnalysisProject(
        id=_new_id(),
        conversation=conversation,
        title=f"Аналитический проект — {conversation.title}",
        version=1,
        document_json={},
    )
    session.add(project)
    session.flush()
    project.document_json = build_project_document(project)
    session.commit()
    return project


def _present(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _formatted_decimal(
    value: object,
    *,
    divisor: Decimal = Decimal("1"),
    suffix: str = "",
    show_plus: bool = False,
) -> str:
    if not _present(value):
        return "—"
    try:
        number = Decimal(str(value).replace(",", ".")) / divisor
    except (InvalidOperation, ValueError):
        return str(value)
    rendered = f"{abs(number):,.2f}".replace(",", "\u00a0").replace(".", ",")
    sign = "−" if number < 0 else "+" if show_plus and number > 0 else ""
    return f"{sign}{rendered}{suffix}"


_REPORT_NUMBER = (
    r"[+-]?(?:\d{1,3}(?:[ \u00a0\u202f]\d{3})+|\d+)(?:[.,]\d+)?"
)
_REPORT_MONEY_RE = re.compile(
    rf"(?<![\w-])(?P<number>{_REPORT_NUMBER})\s*тыс\.?\s*руб\.?",
    re.IGNORECASE,
)
_REPORT_PERCENT_RE = re.compile(
    rf"(?<![\w-])(?P<number>{_REPORT_NUMBER})\s*%",
)
_REPORT_LONG_DECIMAL_RE = re.compile(
    r"(?<![\w-])(?P<number>[+-]?\d+[.,]\d{3,})(?![\w-])"
)


def _decimal_from_text(value: str) -> Decimal | None:
    normalized = (
        value.replace(" ", "")
        .replace("\u00a0", "")
        .replace("\u202f", "")
        .replace(",", ".")
    )
    try:
        return Decimal(normalized)
    except InvalidOperation:
        return None


def format_report_text_numbers(value: object) -> str:
    """Round report prose without modifying the stored analytical evidence."""
    text = str(value if value is not None else "")

    def replace_money(match: re.Match[str]) -> str:
        raw = match.group("number")
        number = _decimal_from_text(raw)
        if number is None:
            return match.group(0)
        magnitude = abs(number)
        if magnitude >= Decimal("1000000"):
            divisor, unit = Decimal("1000000"), "млрд руб."
        elif magnitude >= Decimal("1000"):
            divisor, unit = Decimal("1000"), "млн руб."
        else:
            divisor, unit = Decimal("1"), "тыс. руб."
        return _formatted_decimal(
            number,
            divisor=divisor,
            suffix=f" {unit}",
            show_plus=raw.lstrip().startswith("+"),
        )

    def replace_percent(match: re.Match[str]) -> str:
        raw = match.group("number")
        return _formatted_decimal(
            raw.replace(" ", "")
            .replace("\u00a0", "")
            .replace("\u202f", ""),
            suffix=" %",
            show_plus=raw.lstrip().startswith("+"),
        )

    def replace_long_decimal(match: re.Match[str]) -> str:
        raw = match.group("number")
        return _formatted_decimal(
            raw,
            show_plus=raw.lstrip().startswith("+"),
        )

    text = _REPORT_MONEY_RE.sub(replace_money, text)
    text = re.sub(r"(руб\.)\.+", r"\1", text, flags=re.IGNORECASE)
    text = _REPORT_PERCENT_RE.sub(replace_percent, text)
    return _REPORT_LONG_DECIMAL_RE.sub(replace_long_decimal, text)


def _money_scale(
    rows: list[dict],
    keys: tuple[str, ...] = ("current", "previous", "difference"),
) -> tuple[Decimal, str]:
    maximum = Decimal("0")
    for row in rows:
        if row.get("unit") != "тыс. руб.":
            continue
        for key in keys:
            if not _present(row.get(key)):
                continue
            try:
                maximum = max(maximum, abs(Decimal(str(row[key]))))
            except (InvalidOperation, ValueError):
                continue
    if maximum >= Decimal("1000000"):
        return Decimal("1000000"), "млрд руб."
    if maximum >= Decimal("1000"):
        return Decimal("1000"), "млн руб."
    return Decimal("1"), "тыс. руб."


def _money_value(
    row: dict,
    key: str,
    *,
    divisor: Decimal,
    show_plus: bool = False,
) -> str:
    return _formatted_decimal(
        row.get(key),
        divisor=divisor,
        show_plus=show_plus,
    )


def _extreme_region(row: dict, key: str) -> str:
    item = row.get(key)
    if not isinstance(item, dict) or not item.get("entity"):
        return "—"
    growth = _formatted_decimal(
        item.get("growth_rate"),
        suffix=" %",
        show_plus=True,
    )
    return f"{item['entity']}: {growth}"


def project_section_table(section: dict) -> list[list[str]]:
    """Build the canonical rounded table shared by Web, AI and DOCX."""
    rows = [
        item for item in (section.get("rows") or []) if isinstance(item, dict)
    ]
    if not rows:
        return []
    if section.get("kind") == "regional_average_growth":
        table_rows = [
            [
                "Отрасль",
                "Регионов",
                "Средний прирост, %",
                "Медианный прирост, %",
                "Взвешенный прирост, %",
                "Минимум",
                "Максимум",
            ]
        ]
        table_rows.extend(
            [
                str(row.get("entity", "")),
                str(row.get("regions_count", "—")),
                _formatted_decimal(
                    row.get("growth_rate"),
                    suffix=" %",
                    show_plus=True,
                ),
                _formatted_decimal(
                    row.get("median_growth_rate"),
                    suffix=" %",
                    show_plus=True,
                ),
                _formatted_decimal(
                    row.get("weighted_growth_rate"),
                    suffix=" %",
                    show_plus=True,
                ),
                _extreme_region(row, "minimum_region"),
                _extreme_region(row, "maximum_region"),
            ]
            for row in rows
        )
        return table_rows

    regional_industry = section.get("kind") in {
        "regional_industry_growth",
        "regional_industry_growth_worst",
    }
    monetary = any(row.get("unit") == "тыс. руб." for row in rows)
    money_divisor, money_unit = (
        _money_scale(rows) if monetary else (Decimal("1"), "")
    )
    columns: list[tuple[str, str]] = [
        ("Регион" if regional_industry else "Объект", "entity")
    ]
    for header, key in (
        (
            f"Текущий период, {money_unit}"
            if monetary
            else "Текущий период",
            "current",
        ),
        (
            f"Предыдущий период, {money_unit}"
            if monetary
            else "Предыдущий период",
            "previous",
        ),
        (
            f"Изменение, {money_unit}" if monetary else "Изменение",
            "difference",
        ),
        ("Рост, %", "growth_rate"),
    ):
        if any(_present(row.get(key)) for row in rows):
            columns.append((header, key))
    if any(_present(row.get("score")) for row in rows):
        columns.append(("Индекс", "score"))
    if any(row.get("low_base") or row.get("flagged") for row in rows):
        columns.append(("Примечание", "status"))

    table_rows = [[header for header, _ in columns]]
    for row in rows:
        values: list[str] = []
        for _, key in columns:
            if key == "entity":
                value = str(row.get("entity", ""))
                if regional_industry:
                    region = row.get("region")
                    industry = str(row.get("industry") or "")
                    prefix = f"{industry} — " if industry else ""
                    if region:
                        value = str(region)
                    elif prefix and value.startswith(prefix):
                        value = value[len(prefix):]
            elif key in {"current", "previous"}:
                value = _money_value(
                    row,
                    key,
                    divisor=money_divisor,
                )
            elif key == "difference":
                value = _money_value(
                    row,
                    key,
                    divisor=money_divisor,
                    show_plus=True,
                )
            elif key == "growth_rate":
                value = _formatted_decimal(
                    row.get(key),
                    suffix=" %",
                    show_plus=True,
                )
            elif key == "score":
                value = _formatted_decimal(row.get(key))
            else:
                notes = []
                if row.get("low_base"):
                    notes.append("Малая база")
                if row.get("flagged"):
                    notes.append("Аномальное значение")
                value = "; ".join(notes) or "—"
            values.append(value)
        table_rows.append(values)
    return table_rows


def project_series_table(series: dict) -> list[list[str]]:
    unit = str(series.get("unit", ""))
    monetary = unit == "тыс. руб."
    points = [
        point
        for point in (series.get("points") or [])
        if isinstance(point, dict)
    ]
    money_divisor, money_unit = (
        _money_scale(
            [
                {"unit": unit, "value": point.get("value")}
                for point in points
            ],
            keys=("value",),
        )
        if monetary
        else (Decimal("1"), "")
    )
    header = (
        f"Значение, {money_unit}"
        if monetary
        else f"Значение, {unit}".rstrip(", ")
    )
    rows = [["Период", header]]
    for point in points:
        rows.append(
            [
                str(point.get("label", "")),
                _formatted_decimal(
                    point.get("value"),
                    divisor=money_divisor,
                ),
            ]
        )
    return rows


def project_series_chart(
    block_id: str,
    series: list[dict],
    *,
    title: str = "Динамика показателя",
) -> dict | None:
    usable = [
        item
        for item in series
        if isinstance(item, dict) and item.get("points")
    ]
    if not usable:
        return None
    period_labels = {
        str(point.get("period_end") or ""): str(
            point.get("label")
            or point.get("period_end")
            or ""
        )
        for item in usable
        for point in item.get("points", [])
        if isinstance(point, dict) and point.get("period_end")
    }
    periods = sorted(period_labels)
    chart_type = "bar" if len(periods) == 1 else "line"
    axis_labels = dict(period_labels)
    period_context = ""
    if len(periods) > 1:
        comparable_labels: list[tuple[str, str]] = []
        for period in periods:
            match = re.match(
                r"^(.*?)(?:\s+|[-–—])(\d{4})$",
                period_labels[period].strip(),
            )
            if match is None:
                comparable_labels = []
                break
            comparable_labels.append(
                (match.group(1).strip(" -–—"), match.group(2))
            )
        if (
            comparable_labels
            and len({prefix for prefix, _year in comparable_labels}) == 1
            and len({year for _prefix, year in comparable_labels})
            == len(comparable_labels)
        ):
            period_context = comparable_labels[0][0]
            axis_labels = {
                period: year
                for period, (_prefix, year) in zip(
                    periods,
                    comparable_labels,
                    strict=True,
                )
            }
    label_parts = [
        [part.strip() for part in str(item.get("label", "")).split(" · ")]
        for item in usable
    ]
    common_length = 0
    if len(usable) > 1:
        shortest = min(len(parts) for parts in label_parts)
        while (
            common_length < shortest - 1
            and all(
                parts[common_length] == label_parts[0][common_length]
                for parts in label_parts
            )
        ):
            common_length += 1

    def short_label(index: int, item: dict) -> str:
        value = (
            " · ".join(label_parts[index][common_length:])
            if common_length
            else str(item.get("label", "Ряд"))
        )
        return re.sub(
            r"\s*[-–—]\s*всего\s*$",
            "",
            value,
            flags=re.IGNORECASE,
        ).strip()

    units = {str(item.get("unit") or "") for item in usable}
    return {
        "id": f"{block_id}-chart",
        "type": "chart",
        "title": title,
        "context": (
            " · ".join(label_parts[0][:common_length])
            if common_length
            else ""
        ),
        "chart_type": chart_type,
        "period_context": period_context,
        "mode": (
            "change"
            if chart_type == "line" and len(usable) > 1
            else "absolute"
        ),
        "unit": next(iter(units)) if len(units) == 1 else "",
        "series": [
            {
                "key": str(item.get("key") or f"series-{index + 1}"),
                "label": short_label(index, item),
                "unit": str(item.get("unit") or ""),
                "points": [
                    {
                        "period_end": str(point.get("period_end") or ""),
                        "label": str(
                            point.get("label")
                            or point.get("period_end")
                            or ""
                        ),
                        "axis_label": axis_labels.get(
                            str(point.get("period_end") or ""),
                            str(
                                point.get("label")
                                or point.get("period_end")
                                or ""
                            ),
                        ),
                        "value": str(point.get("value") or "0"),
                    }
                    for point in item.get("points", [])
                    if isinstance(point, dict)
                ],
            }
            for index, item in enumerate(usable)
        ],
        "editable": False,
    }


def build_project_document(project: AnalysisProject) -> dict:
    blocks: list[dict] = [
        {
            "id": "project-title",
            "type": "paragraph",
            "text": project.title,
            "style": "Heading 1",
            "editable": True,
        }
    ]
    seen_table_signatures: set[tuple[tuple[str, ...], ...]] = set()

    def append_unique_table(block_id: str, rows: list[list[str]]) -> bool:
        if not rows:
            return False
        signature = tuple(tuple(str(cell) for cell in row) for row in rows)
        if signature in seen_table_signatures:
            return False
        seen_table_signatures.add(signature)
        blocks.append(
            {
                "id": block_id,
                "type": "table",
                "rows": rows,
                "editable": False,
            }
        )
        return True

    for project_block in sorted(project.blocks, key=lambda item: item.position):
        snapshot = project_block.snapshot_json or {}
        result = snapshot.get("result") or {}
        blocks.extend(
            [
                {
                    "id": f"{project_block.id}-question",
                    "type": "paragraph",
                    "text": snapshot.get("question", "Аналитический вывод"),
                    "style": "Heading 2",
                    "editable": True,
                },
                {
                    "id": f"{project_block.id}-answer",
                    "type": "paragraph",
                    "text": format_report_text_numbers(
                        result.get("answer", "")
                    ),
                    "style": "Normal",
                    "editable": True,
                },
            ]
        )
        research_plan = result.get("research_plan") or {}
        if research_plan:
            questions = research_plan.get("research_questions") or []
            plan_text = str(research_plan.get("objective") or "").strip()
            if questions:
                plan_text += (
                    ("\n" if plan_text else "")
                    + "\n".join(f"• {item}" for item in questions)
                )
            blocks.extend(
                [
                    {
                        "id": f"{project_block.id}-research-plan-title",
                        "type": "paragraph",
                        "text": "План исследования",
                        "style": "Heading 3",
                        "editable": True,
                    },
                    {
                        "id": f"{project_block.id}-research-plan",
                        "type": "paragraph",
                        "text": plan_text,
                        "style": "Normal",
                        "editable": True,
                    },
                ]
            )
        claims = result.get("claims") or []
        if claims:
            blocks.extend(
                [
                    {
                        "id": f"{project_block.id}-claims-title",
                        "type": "paragraph",
                        "text": "Проверенные выводы",
                        "style": "Heading 3",
                        "editable": True,
                    },
                    {
                        "id": f"{project_block.id}-claims",
                        "type": "paragraph",
                        "text": "\n".join(
                            "• "
                            + format_report_text_numbers(
                                item.get("text", "")
                            )
                            for item in claims
                        ),
                        "style": "Normal",
                        "editable": True,
                    },
                ]
            )
        hypotheses = result.get("hypotheses") or []
        if hypotheses:
            blocks.extend(
                [
                    {
                        "id": f"{project_block.id}-hypotheses-title",
                        "type": "paragraph",
                        "text": "Гипотезы",
                        "style": "Heading 3",
                        "editable": True,
                    },
                    {
                        "id": f"{project_block.id}-hypotheses",
                        "type": "paragraph",
                        "text": "\n".join(
                            "• Гипотеза: "
                            + format_report_text_numbers(
                                item.get("text", "")
                            )
                            for item in hypotheses
                        ),
                        "style": "Normal",
                        "editable": True,
                    },
                ]
            )
        series = result.get("series") or []
        chart = project_series_chart(
            project_block.id,
            series,
            title="Динамика по выбранным срезам",
        )
        if chart:
            blocks.append(chart)
        for section_index, section in enumerate(
            result.get("analysis_sections") or []
        ):
            table_rows = project_section_table(section)
            if table_rows:
                signature = tuple(
                    tuple(str(cell) for cell in row) for row in table_rows
                )
                if signature in seen_table_signatures:
                    continue
            blocks.append(
                {
                    "id": f"{project_block.id}-section-{section_index}-title",
                    "type": "paragraph",
                    "text": section.get("title", "Результаты"),
                    "style": "Heading 3",
                    "editable": True,
                }
            )
            if table_rows:
                append_unique_table(
                    f"{project_block.id}-section-{section_index}-table",
                    table_rows,
                )
            methodology = str(section.get("methodology", "")).strip()
            if methodology:
                blocks.append(
                    {
                        "id": (
                            f"{project_block.id}-section-"
                            f"{section_index}-methodology"
                        ),
                        "type": "paragraph",
                        "text": f"Методика: {methodology}",
                        "style": "Normal",
                        "editable": True,
                    }
                )
        for series_index, item in enumerate(series):
            table_rows = project_series_table(item)
            append_unique_table(
                f"{project_block.id}-series-{series_index}",
                table_rows,
            )
    return {
        "version": project.version,
        "presentation_version": PROJECT_PRESENTATION_VERSION,
        "blocks": blocks,
        "paragraph_count": sum(
            1 for item in blocks if item["type"] == "paragraph"
        ),
        "table_count": sum(1 for item in blocks if item["type"] == "table"),
        "chart_count": sum(1 for item in blocks if item["type"] == "chart"),
    }


def _rebuilt_project_document(project: AnalysisProject) -> dict:
    previous = project.document_json or {}
    edited_text = {
        item["id"]: item["text"]
        for item in previous.get("blocks", [])
        if item.get("type") == "paragraph"
        and not str(item.get("id", "")).startswith("project-executive-")
    }
    rebuilt = build_project_document(project)
    for block in rebuilt["blocks"]:
        if block["type"] == "paragraph" and block["id"] in edited_text:
            if not str(block["id"]).endswith("-claims"):
                block["text"] = format_report_text_numbers(
                    edited_text[block["id"]]
                )
    return rebuilt


def rebuild_project_document(project: AnalysisProject) -> dict:
    """Refresh calculated presentation while preserving user prose edits."""
    return _rebuilt_project_document(project)


def project_to_dict(project: AnalysisProject | None) -> dict | None:
    if project is None:
        return None
    document = project.document_json or build_project_document(project)
    return {
        "id": project.id,
        "conversation_id": project.conversation_id,
        "title": project.title,
        "version": project.version,
        "blocks_count": len(project.blocks),
        "blocks": [
            {
                "id": item.id,
                "turn_id": item.turn_id,
                "position": item.position,
                "question": (item.snapshot_json or {}).get(
                    "question", "Аналитический результат"
                ),
            }
            for item in sorted(
                project.blocks, key=lambda block: block.position
            )
        ],
        "document": document,
        "download_url": (
            f"/api/analysis-projects/{project.id}/download/docx"
            if project.output_docx_path
            else None
        ),
        "created_at": project.created_at.isoformat(),
        "updated_at": project.updated_at.isoformat(),
    }


def add_project_block(
    session: Session,
    conversation: Conversation,
    turn: ConversationTurn,
) -> dict:
    if turn.conversation_id != conversation.id or turn.status != "ready":
        raise ValueError("В проект можно добавить только готовый ход этого чата")
    project = ensure_project(session, conversation)
    existing = session.scalar(
        select(AnalysisProjectBlock).where(
            AnalysisProjectBlock.project_id == project.id,
            AnalysisProjectBlock.turn_id == turn.id,
        )
    )
    if existing is None:
        position = (
            session.scalar(
                select(func.max(AnalysisProjectBlock.position)).where(
                    AnalysisProjectBlock.project_id == project.id
                )
            )
            or -1
        ) + 1
        existing = AnalysisProjectBlock(
            id=_new_id(),
            project=project,
            turn_id=turn.id,
            position=position,
            snapshot_json={
                "question": turn.content,
                "plan": turn.plan_json,
                "result": turn.result_json,
            },
        )
        session.add(existing)
        session.flush()
        project.version += 1
        project.updated_at = _now()
        project.document_json = _rebuilt_project_document(project)
        project.output_docx_path = None
        session.commit()
    return {
        "block_id": existing.id,
        "added": True,
        "project": project_to_dict(project),
    }


def remove_project_block(
    session: Session, project: AnalysisProject, block: AnalysisProjectBlock
) -> dict:
    if block in project.blocks:
        project.blocks.remove(block)
    session.delete(block)
    session.flush()
    for index, item in enumerate(
        sorted(project.blocks, key=lambda current: current.position)
    ):
        item.position = index
    project.version += 1
    project.updated_at = _now()
    project.document_json = _rebuilt_project_document(project)
    project.output_docx_path = None
    session.commit()
    return project_to_dict(project) or {}


def update_project(
    session: Session,
    project: AnalysisProject,
    *,
    base_version: int,
    title: str | None,
    ordered_block_ids: list[str] | None,
    changed_text_blocks: dict[str, str],
) -> dict:
    if project.version != base_version:
        raise RuntimeError("version_conflict")
    if title:
        project.title = title.strip()
    if ordered_block_ids is not None:
        current = {item.id: item for item in project.blocks}
        if set(ordered_block_ids) != set(current):
            raise ValueError("Список блоков проекта неполон")
        for position, block_id in enumerate(ordered_block_ids):
            current[block_id].position = position
    document = (
        _rebuilt_project_document(project)
        if ordered_block_ids is not None
        else project.document_json or build_project_document(project)
    )
    if changed_text_blocks:
        by_id = {item["id"]: item for item in document["blocks"]}
        for block_id, candidate in changed_text_blocks.items():
            block = by_id.get(block_id)
            if block is None or block["type"] != "paragraph":
                raise ValueError(f"Текстовый блок {block_id} не найден")
            validate_text_block_edit(block["text"], candidate)
            block["text"] = candidate
    project.version += 1
    document["version"] = project.version
    project.document_json = document
    project.output_docx_path = None
    project.updated_at = _now()
    session.commit()
    return project_to_dict(project) or {}


def _render_project_chart(
    block: dict,
    target: Path,
    settings: Settings,
) -> None:
    cache_dir = settings.app_data_dir / "matplotlib-cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("MPLCONFIGDIR", str(cache_dir))
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt

    series = [
        item
        for item in block.get("series", [])
        if isinstance(item, dict) and item.get("points")
    ]
    periods = sorted(
        {
            str(point.get("period_end") or "")
            for item in series
            for point in item.get("points", [])
            if isinstance(point, dict) and point.get("period_end")
        }
    )
    labels_by_period = {
        str(point.get("period_end") or ""): str(
            point.get("axis_label")
            or point.get("label")
            or point.get("period_end")
            or ""
        )
        for item in series
        for point in item.get("points", [])
        if isinstance(point, dict)
    }
    chart_type = str(
        block.get("chart_type")
        or ("bar" if len(periods) == 1 else "line")
    )
    relative = (
        chart_type == "line"
        and block.get("mode") == "change"
        and len(series) > 1
    )
    colors = [
        "#2563EB",
        "#DC2626",
        "#059669",
        "#7C3AED",
        "#D97706",
        "#0891B2",
        "#DB2777",
        "#4F46E5",
    ]
    figure, axis = plt.subplots(figsize=(10, 5.2), dpi=150)
    for index, item in enumerate(series):
        by_period = {
            str(point.get("period_end") or ""): Decimal(
                str(point.get("value") or "0")
            )
            for point in item.get("points", [])
            if isinstance(point, dict)
        }
        values = [by_period.get(period) for period in periods]
        first = next((value for value in values if value is not None), None)
        rendered: list[float] = []
        for value in values:
            if value is None:
                rendered.append(float("nan"))
            elif relative:
                rendered.append(
                    float((value / first - 1) * 100)
                    if first not in {None, Decimal("0")}
                    else 0.0
                )
            elif item.get("unit") == "тыс. руб.":
                rendered.append(float(value / Decimal("1000000")))
            else:
                rendered.append(float(value))
        label = str(item.get("label") or f"Ряд {index + 1}")
        if len(label) > 58:
            label = label[:55].rstrip() + "…"
        label = textwrap.fill(label, width=32)
        if chart_type == "bar":
            series_count = max(1, len(series))
            bar_width = min(0.72 / series_count, 0.22)
            offsets = [
                position
                - (bar_width * series_count) / 2
                + bar_width / 2
                + index * bar_width
                for position in range(len(periods))
            ]
            axis.bar(
                offsets,
                rendered,
                width=bar_width,
                color=colors[index % len(colors)],
                label=label,
            )
        else:
            axis.plot(
                range(len(periods)),
                rendered,
                color=colors[index % len(colors)],
                linewidth=2.2,
                marker="o",
                markersize=4,
                label=label,
            )
    axis.set_xticks(
        range(len(periods)),
        [labels_by_period.get(period, period) for period in periods],
    )
    axis.tick_params(axis="x", labelsize=8, pad=7)
    if periods:
        axis.set_xlim(-0.48, len(periods) - 0.52)
    period_context = str(block.get("period_context") or "").strip()
    if period_context:
        axis.set_xlabel(period_context, fontsize=8, color="#64748B")
    axis.set_ylabel(
        "% к первому периоду"
        if relative
        else "млрд руб."
        if block.get("unit") == "тыс. руб."
        else str(block.get("unit") or "Значение")
    )
    axis.grid(axis="y", color="#D9E2EC", linewidth=0.8)
    axis.spines[["top", "right"]].set_visible(False)
    if relative:
        axis.axhline(0, color="#64748B", linewidth=0.9)
    if len(series) > 1:
        axis.legend(
            loc="upper center",
            bbox_to_anchor=(0.5, -0.2),
            ncol=2,
            frameon=False,
            fontsize=7,
            handlelength=1.8,
            columnspacing=1.4,
        )
    figure.tight_layout(pad=1.4)
    target.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(
        target,
        bbox_inches="tight",
        pad_inches=0.22,
        facecolor="white",
    )
    plt.close(figure)


def _set_docx_typeface(target, name: str = "Aptos") -> None:
    target.font.name = name
    target._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _set_docx_style_typeface(style, name: str = "Aptos") -> None:
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _set_paragraph_shading(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def _set_paragraph_borders(
    paragraph,
    *,
    color: str,
    sides: tuple[str, ...] = ("bottom",),
    size: int = 6,
    space: int = 4,
) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    for side in sides:
        tag = qn(f"w:{side}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{side}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), str(space))
        border.set(qn("w:color"), color)


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge, color, size in (
        ("top", _DOCX_BORDER, 7),
        ("left", _DOCX_BORDER, 7),
        ("bottom", _DOCX_BORDER, 7),
        ("right", _DOCX_BORDER, 7),
        ("insideH", "E5EBF1", 5),
        ("insideV", "E5EBF1", 5),
    ):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _keep_table_row_together(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = properties.find(qn("w:tblHeader"))
    if repeat is None:
        repeat = OxmlElement("w:tblHeader")
        properties.append(repeat)
    repeat.set(qn("w:val"), "true")


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, result, end))
    _set_docx_typeface(run)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(_DOCX_MUTED)


def _configure_project_docx(document: Document, project: AnalysisProject) -> None:
    properties = document.core_properties
    properties.title = project.title
    properties.subject = "Проверяемый аналитический отчёт 1-МСП"
    properties.author = "ИИ-ВЕКТОР"
    properties.keywords = "1-МСП, аналитика, проверяемые данные"

    normal = document.styles["Normal"]
    _set_docx_style_typeface(normal)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(_DOCX_TEXT)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)

    for style_name, size, color, before, after in (
        ("Heading 1", 24, _DOCX_NAVY, 0, 10),
        ("Heading 2", 16, _DOCX_NAVY, 18, 7),
        ("Heading 3", 11.5, _DOCX_ACCENT, 11, 5),
    ):
        style = document.styles[style_name]
        _set_docx_style_typeface(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 1"].paragraph_format.line_spacing = 1.05
    document.styles["Heading 2"].paragraph_format.line_spacing = 1.08

    bullet = document.styles["List Bullet"]
    _set_docx_style_typeface(bullet)
    bullet.font.size = Pt(10)
    bullet.font.color.rgb = RGBColor.from_string(_DOCX_TEXT)
    bullet.paragraph_format.left_indent = Cm(0.65)
    bullet.paragraph_format.first_line_indent = Cm(-0.25)
    bullet.paragraph_format.space_after = Pt(4)


def _configure_project_section(section, *, landscape: bool) -> None:
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.55)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = (
            section.page_height,
            section.page_width,
        )

    header = section.header.paragraphs[0]
    header.text = "ИИ-ВЕКТОР  /  АНАЛИТИЧЕСКИЙ ОТЧЁТ 1-МСП"
    header.paragraph_format.space_after = Pt(4)
    _set_paragraph_borders(
        header,
        color=_DOCX_ACCENT,
        sides=("bottom",),
        size=12,
        space=3,
    )
    for run in header.runs:
        _set_docx_typeface(run)
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(_DOCX_ACCENT)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(4)
    footer.add_run("ИИ-ВЕКТОР · проверяемые данные · стр. ")
    _add_page_number(footer)
    for run in footer.runs[:-1]:
        _set_docx_typeface(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(_DOCX_MUTED)


def _add_report_accent(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = Pt(1)
    _set_paragraph_borders(
        paragraph,
        color=_DOCX_ACCENT,
        sides=("bottom",),
        size=26,
        space=0,
    )


def _callout_colors(block_id: str) -> tuple[str, str] | None:
    if block_id == "project-executive-summary":
        return _DOCX_ACCENT_SOFT, _DOCX_ACCENT
    if block_id == "project-executive-findings" or block_id.endswith("-claims"):
        return _DOCX_SUCCESS_SOFT, _DOCX_SUCCESS
    if block_id == "project-executive-caveats" or block_id.endswith(
        ("-hypotheses", "-research-plan")
    ):
        return _DOCX_WARNING_SOFT, _DOCX_WARNING
    return None


def _style_project_paragraph(paragraph, block_id: str, style: str) -> None:
    if style == "Heading 2":
        _set_paragraph_borders(
            paragraph,
            color=_DOCX_BORDER,
            sides=("bottom",),
            size=5,
            space=5,
        )
    if block_id == "project-subtitle":
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(9)
        for run in paragraph.runs:
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string(_DOCX_MUTED)
    elif block_id.endswith("-editorial-narrative"):
        paragraph.paragraph_format.left_indent = Cm(0.25)
        paragraph.paragraph_format.right_indent = Cm(0.15)
        _set_paragraph_borders(
            paragraph,
            color="B9CDF5",
            sides=("left",),
            size=10,
            space=6,
        )
    elif block_id.endswith("-methodology"):
        _set_paragraph_shading(paragraph, _DOCX_ROW_ALT)
        for run in paragraph.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(_DOCX_MUTED)

    callout = _callout_colors(block_id)
    if callout is not None:
        fill, accent = callout
        _set_paragraph_shading(paragraph, fill)
        _set_paragraph_borders(
            paragraph,
            color=accent,
            sides=("left",),
            size=16,
            space=7,
        )
        paragraph.paragraph_format.left_indent = Cm(0.28)
        paragraph.paragraph_format.right_indent = Cm(0.2)
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)


def _add_project_text_block(document: Document, block: dict) -> list:
    block_id = str(block.get("id") or "")
    style = str(block.get("style") or "Normal")
    word_style = style if style in {
        "Heading 1",
        "Heading 2",
        "Heading 3",
    } else None
    text = str(block.get("text") or "")
    lines = text.splitlines()
    has_bullets = any(line.strip().startswith("•") for line in lines)
    paragraphs = []
    if not has_bullets:
        paragraphs.append(document.add_paragraph(text, style=word_style))
    else:
        for line in lines:
            value = line.strip()
            if not value:
                continue
            if value.startswith("•"):
                paragraph = document.add_paragraph(
                    value.lstrip("•").strip(),
                    style="List Bullet",
                )
            else:
                paragraph = document.add_paragraph(value, style=word_style)
            paragraphs.append(paragraph)
        if not paragraphs:
            paragraphs.append(document.add_paragraph("", style=word_style))
    for paragraph in paragraphs:
        _style_project_paragraph(paragraph, block_id, style)
    return paragraphs


def _add_project_metadata(
    document: Document,
    project: AnalysisProject,
    document_model: dict,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.left_indent = Cm(0.18)
    paragraph.paragraph_format.right_indent = Cm(0.18)
    _set_paragraph_shading(paragraph, "F4F7FB")
    _set_paragraph_borders(
        paragraph,
        color=_DOCX_BORDER,
        sides=("top", "bottom"),
        size=5,
        space=5,
    )
    values = (
        f"ВЕРСИЯ {project.version}",
        f"ПРОВЕРЕННЫХ РЕЗУЛЬТАТОВ {len(project.blocks)}",
        f"ТАБЛИЦ {document_model.get('table_count', 0)}",
        f"ГРАФИКОВ {document_model.get('chart_count', 0)}",
    )
    for index, value in enumerate(values):
        if index:
            separator = paragraph.add_run("    •    ")
            separator.font.color.rgb = RGBColor.from_string("A9B7C6")
        run = paragraph.add_run(value)
        _set_docx_typeface(run)
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(_DOCX_MUTED)


def _add_project_table(document: Document, block: dict) -> None:
    rows = block.get("rows") or [[]]
    width = max((len(row) for row in rows), default=1)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders(table)
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        _keep_table_row_together(row)
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(
                cell,
                top=90,
                start=105,
                bottom=90,
                end=105,
            )
            if row_index == 0:
                _set_cell_shading(cell, _DOCX_NAVY)
            elif row_index % 2 == 0:
                _set_cell_shading(cell, _DOCX_ROW_ALT)
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if column_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    _set_docx_typeface(run)
                    run.font.size = Pt(7.5 if width >= 7 else 8.5)
                    run.bold = row_index == 0 or (
                        column_index == 0 and row_index > 0
                    )
                    run.font.color.rgb = RGBColor.from_string(
                        "FFFFFF" if row_index == 0 else _DOCX_TEXT
                    )
        if row_index == 0:
            _repeat_table_header(row)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def _add_project_chart(
    document: Document,
    block: dict,
    chart_path: Path,
    *,
    landscape: bool,
) -> None:
    eyebrow = document.add_paragraph("ГРАФИК")
    eyebrow.paragraph_format.space_before = Pt(9)
    eyebrow.paragraph_format.space_after = Pt(2)
    for run in eyebrow.runs:
        _set_docx_typeface(run)
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(_DOCX_ACCENT)

    title = document.add_paragraph(
        str(block.get("title") or "Динамика"),
        style="Heading 3",
    )
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    context = str(block.get("context") or "").strip()
    if context:
        context_paragraph = document.add_paragraph(context)
        context_paragraph.paragraph_format.space_after = Pt(5)
        for run in context_paragraph.runs:
            _set_docx_typeface(run)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(_DOCX_MUTED)

    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(3)
    picture.paragraph_format.space_after = Pt(10)
    _set_paragraph_shading(picture, "F8FAFD")
    _set_paragraph_borders(
        picture,
        color=_DOCX_BORDER,
        sides=("top", "left", "bottom", "right"),
        size=5,
        space=7,
    )
    picture.add_run().add_picture(
        str(chart_path),
        width=Inches(9.25 if landscape else 6.35),
    )


def write_project_docx(project: AnalysisProject, settings: Settings) -> Path:
    document_model = project.document_json or build_project_document(project)
    if (
        document_model.get("presentation_version")
        != PROJECT_PRESENTATION_VERSION
    ):
        document_model = rebuild_project_document(project)
        project.document_json = document_model
    document = Document()
    _configure_project_docx(document, project)
    max_columns = max(
        (
            max((len(row) for row in block.get("rows", [])), default=0)
            for block in document_model["blocks"]
            if block["type"] == "table"
        ),
        default=0,
    )
    landscape = max_columns >= 7
    for section in document.sections:
        _configure_project_section(section, landscape=landscape)
    _add_report_accent(document)
    chart_paths: list[Path] = []
    metadata_added = False
    for block in document_model["blocks"]:
        if block["type"] == "paragraph":
            _add_project_text_block(document, block)
            block_id = str(block.get("id") or "")
            if block_id == "project-subtitle":
                _add_project_metadata(document, project, document_model)
                metadata_added = True
            elif block_id == "project-title" and not any(
                str(item.get("id") or "") == "project-subtitle"
                for item in document_model["blocks"]
            ):
                _add_project_metadata(document, project, document_model)
                metadata_added = True
        elif block["type"] == "table":
            if not metadata_added:
                _add_project_metadata(document, project, document_model)
                metadata_added = True
            _add_project_table(document, block)
        elif block["type"] == "chart":
            if not metadata_added:
                _add_project_metadata(document, project, document_model)
                metadata_added = True
            chart_path = (
                settings.app_data_dir
                / "reports"
                / f".project-chart-{uuid.uuid4().hex}.png"
            )
            _render_project_chart(block, chart_path, settings)
            chart_paths.append(chart_path)
            _add_project_chart(
                document,
                block,
                chart_path,
                landscape=landscape,
            )
    if not metadata_added:
        _add_project_metadata(document, project, document_model)
    target = (
        settings.app_data_dir
        / "reports"
        / f"analytical_project_{project.id}_v{project.version}.docx"
    )
    temporary = target.with_name(f".{target.name}.{uuid.uuid4().hex}.tmp")
    try:
        document.save(temporary)
        os.replace(temporary, target)
    finally:
        for chart_path in chart_paths:
            chart_path.unlink(missing_ok=True)
    project.output_docx_path = str(target)
    return target


def scenario_to_dict(scenario: Scenario) -> dict:
    return {
        "id": scenario.id,
        "title": scenario.title,
        "source_conversation_id": scenario.source_conversation_id,
        "steps": scenario.steps_json or [],
        "created_at": scenario.created_at.isoformat(),
        "updated_at": scenario.updated_at.isoformat(),
    }


def scenario_steps(conversation: Conversation) -> list[dict]:
    steps = []
    for turn in conversation.turns:
        if turn.status != "ready" or not turn.plan_json:
            continue
        plan_template = dict(turn.plan_json)
        plan_template["periods"] = []
        plan_template["period_labels"] = []
        plan_template["period_modes"] = []
        years = sorted(
            {
                int(period[:4])
                for period in (turn.plan_json or {}).get("periods", [])
                if isinstance(period, str)
                and re.fullmatch(r"20\d{2}-\d{2}-\d{2}", period)
            }
        )
        question_template = turn.content
        if years:
            latest = max(years)
            for year in sorted(years, reverse=True):
                offset = latest - year
                placeholder = (
                    "{{latest_year}}"
                    if offset == 0
                    else f"{{{{latest_year_minus_{offset}}}}}"
                )
                question_template = re.sub(
                    rf"\b{year}\b", placeholder, question_template
                )
        steps.append(
            {
                "question": question_template,
                "plan_template": plan_template,
                "period_strategy": "latest_same_slice",
            }
        )
    return steps
