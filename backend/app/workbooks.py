from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Protocol

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


DATE_RE = re.compile(r"(?<!\d)(\d{2})[./](\d{2})[./](\d{4})(?!\d)")
FORMULA_RE = re.compile(r"(?:стр\\.?\\s*)?\\d{4}\\s*[=>]+\\s*(.+)", re.IGNORECASE)


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = text.replace("ё", "е").replace("Ё", "Е")
    text = re.sub(r"[-‐‑‒–—]", "-", text)
    return re.sub(r"\s+", " ", text).strip().casefold()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decimal_value(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    if isinstance(value, str):
        candidate = value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
        try:
            return Decimal(candidate)
        except InvalidOperation:
            return None
    return None


class SheetReader(Protocol):
    title: str
    max_row: int
    max_column: int

    def value(self, row: int, column: int) -> Any: ...

    def coordinate(self, row: int, column: int) -> str: ...


class WorkbookReader(Protocol):
    path: Path

    @property
    def sheet_names(self) -> list[str]: ...

    def sheet(self, name: str) -> SheetReader: ...

    def close(self) -> None: ...


class XlsxSheet:
    def __init__(self, worksheet):
        self._sheet = worksheet
        self.title = worksheet.title
        self.max_row = worksheet.max_row
        self.max_column = worksheet.max_column

    def value(self, row: int, column: int) -> Any:
        return self._sheet.cell(row, column).value

    def coordinate(self, row: int, column: int) -> str:
        return f"{get_column_letter(column)}{row}"


class XlsxReader:
    def __init__(self, path: Path, *, extension_mismatch: bool = False):
        self.path = path
        self._handle = path.open("rb") if extension_mismatch else None
        # Random access is required for schema coordinates and evidence. In
        # openpyxl read-only mode every isolated cell lookup reparses XML and is
        # dramatically slower for this workload.
        self._book = load_workbook(
            self._handle or path, data_only=True, read_only=False
        )

    @property
    def sheet_names(self) -> list[str]:
        return self._book.sheetnames

    def sheet(self, name: str) -> XlsxSheet:
        return XlsxSheet(self._book[name])

    def close(self) -> None:
        self._book.close()
        if self._handle is not None:
            self._handle.close()


class XlsSheet:
    def __init__(self, worksheet):
        self._sheet = worksheet
        self.title = worksheet.name
        self.max_row = worksheet.nrows
        self.max_column = worksheet.ncols

    def value(self, row: int, column: int) -> Any:
        if row < 1 or column < 1 or row > self.max_row or column > self.max_column:
            return None
        return self._sheet.cell_value(row - 1, column - 1)

    def coordinate(self, row: int, column: int) -> str:
        return f"{get_column_letter(column)}{row}"


class XlsReader:
    def __init__(self, path: Path):
        self.path = path
        self._book = xlrd.open_workbook(path, on_demand=True)

    @property
    def sheet_names(self) -> list[str]:
        return self._book.sheet_names()

    def sheet(self, name: str) -> XlsSheet:
        return XlsSheet(self._book.sheet_by_name(name))

    def close(self) -> None:
        self._book.release_resources()


def open_reader(path: Path) -> WorkbookReader:
    suffix = path.suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        return XlsxReader(path)
    if suffix == ".xls":
        # A few historical files carry the legacy extension but contain an
        # OOXML ZIP package. Detect the format from its signature.
        with path.open("rb") as handle:
            if handle.read(4).startswith(b"PK"):
                return XlsxReader(path, extension_mismatch=True)
        return XlsReader(path)
    raise ValueError(f"Неподдерживаемый формат: {path.suffix}")


def find_sheet(reader: WorkbookReader, wanted: str) -> str | None:
    wanted_normalized = normalize_text(wanted)
    for name in reader.sheet_names:
        normalized = normalize_text(name)
        if normalized == wanted_normalized or wanted_normalized in normalized:
            return name
    return None


def classify_workbook(path: Path, sheet_names: list[str]) -> str:
    filename = normalize_text(path.stem)
    sheet_text = " ".join(normalize_text(name) for name in sheet_names)
    if path.name.startswith("~$"):
        return "service_file"
    if "коды" in filename:
        return "codes_dictionary"
    if "поступлен" in filename and "1-мсп" not in filename:
        return "national_receipts"
    if "автомсп" in filename or path.suffix.casefold() == ".xlsm":
        return "controller"
    if "аналитика" in filename:
        return "analytical_auxiliary"
    if "раздел 1" in filename or "раздел1" in filename:
        return "regional_section_i"
    if "раздел 2" in filename or "раздел2" in filename:
        return "regional_section_ii"
    if "в разрезе субъектов" in filename and "раздел 1" in filename:
        return "regional_section_i"
    if "в разрезе субъектов" in filename and "раздел 2" in filename:
        return "regional_section_ii"
    if "свод по россии" in filename:
        return "msp_summary"
    if "1-мсп" in filename and (
        "раздел i" in sheet_text
        or "раздел 1" in sheet_text
        or "титульный лист" in sheet_text
    ):
        return "msp_summary"
    if "раздел i" in sheet_text and "раздел ii" in sheet_text:
        return "msp_summary"
    return "unknown"


def extract_as_of_date(reader: WorkbookReader) -> tuple[date | None, list[str]]:
    warnings: list[str] = []
    candidates: list[tuple[str, date]] = []
    preferred: list[date] = []
    sheet_order = sorted(
        reader.sheet_names,
        key=lambda name: 0 if "титул" in normalize_text(name) else 1,
    )
    for name in sheet_order[:4]:
        sheet = reader.sheet(name)
        for row in range(1, min(sheet.max_row, 25) + 1):
            for column in range(1, min(sheet.max_column, 12) + 1):
                raw = sheet.value(row, column)
                if not isinstance(raw, str):
                    continue
                for match in DATE_RE.finditer(raw):
                    day, month, year = map(int, match.groups())
                    try:
                        parsed = date(year, month, day)
                    except ValueError:
                        continue
                    candidates.append((raw, parsed))
                    if "по состоянию" in normalize_text(raw):
                        preferred.append(parsed)
    unique_preferred = sorted(set(preferred))
    if len(unique_preferred) == 1:
        return unique_preferred[0], warnings
    if len(unique_preferred) > 1:
        warnings.append(
            "В книге найдено несколько дат «по состоянию на»: "
            + ", ".join(item.isoformat() for item in unique_preferred)
        )
        return None, warnings
    warnings.append("Дата «по состоянию на» внутри книги не найдена")
    return None, warnings


def period_from_as_of(as_of_date: date) -> tuple[date, date, int, int]:
    period_end = as_of_date - timedelta(days=1)
    return date(period_end.year, 1, 1), period_end, period_end.year, period_end.month


@dataclass(frozen=True)
class HeaderSchema:
    header_number_row: int
    code_column: int
    metric_columns: dict[int, int]
    raw_headers: dict[int, str]
    numeric_count: int
    header_hash: str
    fingerprint: str


def recognize_header(sheet: SheetReader, form_name: str = "1-МСП") -> HeaderSchema:
    best_row = 0
    best_mapping: dict[int, int] = {}
    for row in range(1, min(sheet.max_row, 30) + 1):
        mapping: dict[int, int] = {}
        for column in range(1, min(sheet.max_column, 80) + 1):
            value = sheet.value(row, column)
            number: int | None = None
            if isinstance(value, int):
                number = value
            elif isinstance(value, float) and value.is_integer():
                number = int(value)
            elif isinstance(value, str) and value.strip().isdigit():
                number = int(value.strip())
            if number is not None and 1 <= number <= 100:
                mapping[number] = column
        if len(mapping) > len(best_mapping) and all(n in mapping for n in range(1, min(5, len(mapping)) + 1)):
            best_row, best_mapping = row, mapping
    if len(best_mapping) < 4:
        raise ValueError("Не найдена строка с номерами граф формы")

    first_metric_column = best_mapping[min(best_mapping)]
    # The service columns to the left of graph 1 are not stable across form
    # editions.  In particular, several 2023/2024 Section II books have an
    # empty column immediately before graph 1, while the row code is one
    # column further left.  Resolve the code column from its values instead of
    # assuming a fixed offset from the first metric.
    code_candidates: list[tuple[int, int]] = []
    for column in range(1, first_metric_column):
        matches = 0
        for row in range(best_row + 1, sheet.max_row + 1):
            value = decimal_value(sheet.value(row, column))
            if (
                value is not None
                and value == value.to_integral_value()
                and 1000 <= int(value) < 3000
            ):
                matches += 1
        code_candidates.append((matches, column))
    best_code_matches, code_column = max(
        code_candidates,
        default=(0, first_metric_column - 1),
    )
    if best_code_matches == 0:
        code_column = first_metric_column - 1
    raw_headers: dict[int, str] = {}
    normalized_headers: list[str] = []
    for report_column, physical_column in sorted(best_mapping.items()):
        parts: list[str] = []
        for row in range(max(1, best_row - 6), best_row):
            value = sheet.value(row, physical_column)
            normalized = normalize_text(value)
            if normalized and normalized not in parts:
                parts.append(normalized)
        raw = " → ".join(parts)
        raw_headers[report_column] = raw
        normalized_headers.append(f"{report_column}:{raw}")
    header_hash = hashlib.sha256("|".join(normalized_headers).encode()).hexdigest()
    fingerprint_source = (
        f"{normalize_text(form_name)}|{sheet.title}|{len(best_mapping)}|"
        f"{best_row}|{code_column}|{header_hash}"
    )
    fingerprint = hashlib.sha256(fingerprint_source.encode()).hexdigest()
    return HeaderSchema(
        header_number_row=best_row,
        code_column=code_column,
        metric_columns=best_mapping,
        raw_headers=raw_headers,
        numeric_count=len(best_mapping),
        header_hash=header_hash,
        fingerprint=fingerprint,
    )


def source_files(root: Path) -> list[Path]:
    return sorted(
        path
        for path in root.rglob("*")
        if path.is_file()
        and path.suffix.casefold() in {".xls", ".xlsx", ".xlsm"}
    )
