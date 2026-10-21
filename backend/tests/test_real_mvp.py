from __future__ import annotations

import asyncio
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
import backend.app.llm as llm_module
from docx import Document

from backend.app.calculations import (
    build_report_data,
    find_receipts_workbook,
    safe_ratio,
)
from backend.app.catalog import (
    METRIC_BY_CODE,
    resolve_header_metrics,
    resolve_metric_code,
    resolve_metric_codes,
)
from backend.app.config import Settings
from backend.app.database import SessionLocal
from backend.app.reporting import generate_report
from backend.app.regional import (
    inspect_regional_workbook,
    read_regional_values,
    territory_key,
    territory_type,
)
from backend.app.query_service import build_query_plan, extract_years
from backend.app.llm import (
    GnivcGigaChatProvider,
    GnivcQwenProvider,
    LLMProviderError,
    OpenAICompatibleProvider,
    configured_provider,
    configured_narrative_provider,
    configured_report_synthesis_provider,
)
from backend.app.planner import (
    _apply_deterministic_constraints,
    _enrich_explicit_territories,
    detect_catalog_target,
)
from backend.app.analytical_agent import (
    _candidate_analyses,
    _resolve_named_industry,
)
from backend.app.schemas import QueryPlan
from backend.app.workbooks import (
    decimal_value,
    normalize_text,
    open_reader,
    recognize_header,
    source_files,
)


ROOT = Path(__file__).resolve().parents[2]
HISTORY = ROOT / "rd" / "1_msp_db"
AUTOMSP = ROOT / "rd" / "AutoMSP"
OLD = AUTOMSP / "53778127_1-МСП.xlsx"
NEW = AUTOMSP / "75638291_1-МСП.xlsx"
CODES = AUTOMSP / "Коды.xlsx"
TEMPLATE = AUTOMSP / "template.docx"
REGIONAL_2026 = (
    HISTORY / "2026" / "01.04.2026" / "70976024_Раздел 1.xlsx"
)
NATIONAL_2022 = (
    HISTORY / "2022" / "01.01.2023" / "32916315_1-МСП.xls"
)
NATIONAL_2023_SEPTEMBER = (
    HISTORY / "2023" / "01.10.2023" / "38120763_1-МСП.xls"
)
REQUIRES_HISTORY = pytest.mark.skipif(
    not HISTORY.is_dir(),
    reason="Для интеграционного теста нужен локальный каталог rd/1_msp_db",
)
REQUIRES_AUTOMSP = pytest.mark.skipif(
    not AUTOMSP.is_dir(),
    reason="Для интеграционного теста нужен локальный каталог rd/AutoMSP",
)


def _report_data() -> dict:
    return build_report_data(
        OLD,
        NEW,
        find_receipts_workbook(AUTOMSP),
        CODES,
        old_year=2025,
        new_year=2026,
        months_count=6,
    )


def _row(data: dict, code: str) -> dict:
    return next(item for item in data["table1"] if item["code"] == code)


@REQUIRES_HISTORY
def test_historical_inventory_matches_brief() -> None:
    files = source_files(HISTORY)
    assert len(files) == 166
    assert sum(path.suffix.casefold() == ".xlsx" for path in files) == 126
    assert sum(path.suffix.casefold() == ".xls" for path in files) == 40


@REQUIRES_HISTORY
def test_regional_workbook_streaming_layout() -> None:
    layout = inspect_regional_workbook(REGIONAL_2026)
    assert layout.as_of_date == date(2026, 4, 1)
    assert layout.numeric_count == 43
    assert layout.metric_columns[10] == 13
    assert "1010" in layout.sheet_names
    assert len(layout.sheet_names) == 377


@REQUIRES_HISTORY
def test_2022_total_receipts_are_resolved_by_header_not_current_graph() -> None:
    reader = open_reader(NATIONAL_2022)
    try:
        sheet = reader.sheet("Раздел I")
        schema = recognize_header(sheet)
        resolution = resolve_header_metrics(schema.raw_headers)
        assert resolution.is_complete
        assert resolution.source_to_code[2] == "receipts_budget_total"
        assert resolution.source_to_code[4] == "receipts_profit_tax_total"

        total_row = next(
            row
            for row in range(
                schema.header_number_row + 1,
                sheet.max_row + 1,
            )
            if decimal_value(sheet.value(row, schema.code_column))
            == Decimal("1010")
        )
        assert (
            decimal_value(
                sheet.value(total_row, schema.metric_columns[2])
            )
            == Decimal("5713921391")
        )
        assert (
            decimal_value(
                sheet.value(total_row, schema.metric_columns[4])
            )
            == Decimal("863149266")
        )
    finally:
        reader.close()


@REQUIRES_HISTORY
def test_section_ii_row_code_skips_an_inserted_service_column() -> None:
    reader = open_reader(NATIONAL_2023_SEPTEMBER)
    try:
        sheet = reader.sheet("Раздел II")
        schema = recognize_header(sheet)
        assert schema.metric_columns[1] == 8
        assert schema.code_column == 6
        assert any(
            decimal_value(sheet.value(row, schema.code_column))
            == Decimal("2010")
            for row in range(
                schema.header_number_row + 1,
                sheet.max_row + 1,
            )
        )
    finally:
        reader.close()


@REQUIRES_HISTORY
def test_2026_regional_header_conflict_uses_country_crosscheck() -> None:
    with SessionLocal() as session:
        values = read_regional_values(
            session,
            period=date(2026, 3, 31),
            section="I",
            row_code=1010,
            report_column=METRIC_BY_CODE[
                "receipts_insurance_distributed"
            ].column,
        )
    country = next(
        item for item in values if item.territory_type == "country"
    )
    assert country.value == Decimal("784485425")
    assert country.source["report_column"] == 39
    assert (
        country.source["mapping_method"]
        == "national_schema_country_crosscheck"
    )
    assert country.source["country_crosscheck"] == {
        "comparable": 43,
        "matching": 43,
    }


@REQUIRES_HISTORY
def test_normalization_and_mislabeled_xlsx_adapter() -> None:
    assert normalize_text("  РАЗДЕЛ\u00a0I — Ёлка ") == "раздел i - елка"
    mislabeled = next(HISTORY.rglob("39186960_1-МСП.xls"))
    reader = open_reader(mislabeled)
    try:
        assert "Раздел I" in reader.sheet_names
        assert "Раздел II" in reader.sheet_names
    finally:
        reader.close()


def test_zero_denominator_is_explicit() -> None:
    assert safe_ratio(Decimal("10"), Decimal("0")) is None


def test_broad_msp_receipts_mean_budget_plus_insurance() -> None:
    plan = build_query_plan(
        "Сравни поступления от МСП за январь-июнь 2025 и 2026", []
    )
    assert plan.metric_codes == [
        "receipts_budget_total",
        "receipts_insurance_total",
    ]


def test_exact_accrued_insurance_metric_beats_receipts_alias() -> None:
    question = "Покажи показатель «Начислено страховых взносов» по России"
    assert resolve_metric_code(question) == "accrued_insurance_total"
    assert resolve_metric_codes(question) == ["accrued_insurance_total"]


@pytest.mark.parametrize(
    ("reference", "expected"),
    [
        ("поступления мсп", "receipts_budget_total"),
        ("поступления от мсп", "receipts_budget_total"),
        ("автоусн", "receipts_autousn"),
        ("акцизы на пиво", "receipts_excise_beer"),
        (
            "receipts_insurance_pension_fixed",
            "receipts_insurance_pension_fixed",
        ),
        (
            "receipts_insurance_medical_fixed",
            "receipts_insurance_medical_fixed",
        ),
    ],
)
def test_quoted_metric_reference_is_exact(
    reference: str,
    expected: str,
) -> None:
    question = (
        f"Какое значение показателя «{reference}» по всем МСП по "
        "Российской Федерации за январь-июнь 2026 года?"
    )
    plan = build_query_plan(question, [date(2026, 6, 30)])
    assert plan.metric_codes == [expected]


@pytest.mark.parametrize(
    "metric_name",
    [
        "Федеральные налоги и сборы",
        "Налог на прибыль организаций",
        "Акцизы на алкогольную продукцию",
        "Транспортный налог",
        "Взносы на медицинское страхование",
        "Единый налоговый платёж физического лица",
        "receipts_unified_tax_payment_individual",
        "графа 50",
        "енп физического лица",
    ],
)
def test_quoted_metric_does_not_become_industry_or_federal_district(
    metric_name: str,
) -> None:
    question = (
        f"Какое значение показателя «{metric_name}» по всем МСП по "
        "Российской Федерации за январь-март 2022 года?"
    )
    plan = build_query_plan(question, [date(2022, 3, 31)])
    with SessionLocal() as session:
        plan = _enrich_explicit_territories(session, question, plan)
        industry = _resolve_named_industry(session, question)
    assert plan.territories == ["Российская Федерация"]
    assert industry is None


def test_equivalent_settings_share_one_qwen_connection_pool() -> None:
    llm_module._PROVIDER_CACHE.clear()
    settings = Settings(
        llm_provider="gnivc_qwen",
        gnivc_api_key="test-key",
        gnivc_query_model="qwen3.6-35b",
        llm_max_retries=0,
    )
    first = configured_provider(settings)
    second = configured_provider(settings.model_copy())
    assert first is second
    assert len(llm_module._PROVIDER_CACHE) == 1
    assert isinstance(first, OpenAICompatibleProvider)
    assert first._client is not None
    assert first._client.max_retries == 0
    llm_module._PROVIDER_CACHE.clear()


def test_budget_receipts_keep_requested_enterprise_group() -> None:
    micro = build_query_plan(
        "Каковы поступления в бюджет для микропредприятий по России "
        "за январь-июнь 2026?",
        [date(2026, 6, 30)],
    )
    without_micro = build_query_plan(
        "Каковы бюджетные поступления малых и средних предприятий без "
        "микро по России за январь-июнь 2026?",
        [date(2026, 6, 30)],
    )
    assert micro.metric_codes == ["receipts_budget_total"]
    assert micro.enterprise_groups == ["micro"]
    assert without_micro.metric_codes == ["receipts_budget_total"]
    assert without_micro.enterprise_groups == [
        "small_and_medium_without_micro"
    ]


def test_catalog_row_numbers_are_not_parsed_as_years() -> None:
    assert extract_years(
        "Покажи ВСЕГО (стр.2010=стр.2020+2028+2033) за 2026 год"
    ) == [2026]


def test_word_outliers_selects_regional_outlier_tools() -> None:
    question = (
        "Какие регионы являются выбросами по уровню и росту НДС "
        "за январь-июнь 2026?"
    )
    plan = build_query_plan(question, [date(2026, 6, 30)])
    assert plan.intent == "outliers"
    assert _candidate_analyses(question, plan) == [
        "regional_rank",
        "outliers",
    ]


@pytest.mark.parametrize(
    "reference",
    [
        "Дальневосточный федеральный округ",
        "Дальневосточный ФО",
        "Дальневосточный-федеральный-округ",
        "Дльневосточный федеральный округ",
    ],
)
def test_explicit_territory_has_one_authoritative_match(
    reference: str,
) -> None:
    question = (
        f"Покажи НДС по всем МСП в территории «{reference}» "
        "за январь-март 2026 года."
    )
    plan = build_query_plan(question, [date(2026, 3, 31)])
    # Simulate an unrelated territory guessed by a language-model planner.
    plan.territories = ["Южный федеральный округ"]
    with SessionLocal() as session:
        plan = _enrich_explicit_territories(session, question, plan)
        industry = _resolve_named_industry(session, question)
    assert plan.territories == ["Дальневосточный федеральный округ"]
    assert industry is None


@pytest.mark.parametrize(
    "reference",
    [
        "в Республике Марий Эл",
        "Республика-Марий-Эл",
        "в Республике Марий-Эл",
    ],
)
def test_semantically_duplicate_territory_spellings_resolve(
    reference: str,
) -> None:
    question = (
        f"Покажи НДС по всем МСП в территории «{reference}» "
        "за январь-март 2026 года."
    )
    plan = build_query_plan(question, [date(2026, 3, 31)])
    with SessionLocal() as session:
        plan = _enrich_explicit_territories(session, question, plan)
    assert [territory_key(item) for item in plan.territories] == [
        territory_key("Республика Марий Эл")
    ]


@pytest.mark.parametrize(
    ("reference", "expected"),
    [
        ("Рспублика Татарстан", "Республика Татарстан"),
        ("г. Мсква", "г. Москва"),
        ("г. Снкт-Петербург", "г. Санкт-Петербург"),
        ("г. Свастополь", "г. Севастополь"),
    ],
)
def test_explicit_territory_typo_uses_semantic_catalog_ranking(
    reference: str,
    expected: str,
) -> None:
    question = (
        f"Покажи НДС по всем МСП в территории «{reference}» "
        "за январь-март 2026 года."
    )
    plan = build_query_plan(question, [date(2026, 3, 31)])
    # A prior planner guess must not survive an authoritative quoted match.
    plan.territories = ["Республика Марий Эл"]
    with SessionLocal() as session:
        plan = _enrich_explicit_territories(session, question, plan)
    assert [territory_key(item) for item in plan.territories] == [
        territory_key(expected)
    ]


def test_receipts_together_with_insurance_mean_derived_total() -> None:
    plan = build_query_plan(
        "а что по поступлениям (вместе со страховыми взносами) "
        "московской области?",
        [date(2026, 6, 30)],
    )
    assert plan.metric_codes == [
        "receipts_budget_total",
        "receipts_insurance_total",
    ]
    assert plan.row_codes == [1010]


def test_metric_replacement_keeps_only_target_metric() -> None:
    plan = build_query_plan(
        "Там же замени НДС на НДФЛ.",
        [date(2026, 3, 31)],
    )
    assert plan.metric_codes == ["receipts_ndfl_total"]


def test_insurance_question_with_different_standalone_quarters() -> None:
    plan = build_query_plan(
        "какая динамика по страховым взносам если сравнить "
        "первый квартал 2026 и второй квартал 2025?",
        [],
    )
    assert plan.metric_codes == ["receipts_insurance_total"]
    assert [period.isoformat() for period in plan.periods] == [
        "2026-03-31",
        "2025-06-30",
    ]
    assert plan.period_modes == [
        "standalone_quarter",
        "standalone_quarter",
    ]
    assert plan.period_labels == ["I квартал 2026", "II квартал 2025"]


def test_report_column_number_resolves_metric() -> None:
    plan = build_query_plan(
        "Сравни графу 10 за январь-июнь 2025 и 2026", []
    )
    assert plan.metric_codes == ["receipts_vat_total"]
    assert [period.isoformat() for period in plan.periods] == [
        "2025-06-30",
        "2026-06-30",
    ]


def test_explicit_vat_wins_over_generic_msp_receipts_phrase() -> None:
    plan = build_query_plan(
        "сравни поступления мсп по калининграду и москве по НДС "
        "за первый квартал 2025 и первый квартал 2026",
        [],
    )
    assert plan.metric_codes == ["receipts_vat_total"]
    assert METRIC_BY_CODE["receipts_vat_total"].short_name == "НДС"


def test_metric_aliases_work_in_russian_cases() -> None:
    assert (
        resolve_metric_code("поступления по спецрежимам")
        == "receipts_special_regimes_total"
    )
    assert (
        resolve_metric_code("динамика по специальным налоговым режимам")
        == "receipts_special_regimes_total"
    )
    assert resolve_metric_code("сравнение по акцизам") == "receipts_excise_total"
    assert (
        resolve_metric_code("поступления по налогу на прибыль")
        == "receipts_profit_tax_total"
    )


def test_catalog_questions_are_not_treated_as_numeric_queries() -> None:
    assert (
        detect_catalog_target(
            "Для какого периода можно сравнить поступления МСП "
            "по Калининграду и Москве по НДС?"
        )
        == "periods"
    )
    assert detect_catalog_target("По каким регионам есть данные?") == "territories"
    assert detect_catalog_target("Какие налоги есть?") == "metrics"
    assert detect_catalog_target("Какие коды ОКВЭД доступны?") == "okved"
    assert detect_catalog_target("Что вообще доступно в данных?") == "overview"
    assert detect_catalog_target("Какие налоги выросли больше всего?") is None
    assert (
        detect_catalog_target(
            "Какие поступления по спецрежимам были в Калининграде "
            "за первый квартал 2026?"
        )
        is None
    )


def test_federal_city_aliases_have_one_territory_key() -> None:
    assert territory_key("г. Москва") == territory_key("город Москва")
    assert territory_key("г.Москва") == territory_key("Москва")


def test_regional_statutory_short_names_have_one_territory_key() -> None:
    pairs = [
        ("Кемеровская область", "Кемеровская область - Кузбасс"),
        ("Ненецкий автономный округ", "Ненецкий АО"),
        ("Республика Адыгея (Адыгея)", "Республика Адыгея"),
        (
            "Ханты-Мансийский автономный округ - Югра",
            "Ханты-Мансийский АО - Югра",
        ),
        ("Чукотский автономный округ", "Чукотский АО"),
        ("Ямало-Ненецкий автономный округ", "Ямало-Ненецкий АО"),
    ]
    for official, short in pairs:
        assert territory_key(official) == territory_key(short)
    assert territory_type("Юго-Западные субъекты") == "territory_group"


def test_explicit_quarters_override_an_incorrect_llm_plan() -> None:
    llm_plan = QueryPlan(
        intent="compare",
        metric_codes=["receipts_insurance_total"],
        periods=[date(2025, 3, 31), date(2026, 3, 31)],
        period_modes=["standalone_quarter", "standalone_quarter"],
    )
    corrected = _apply_deterministic_constraints(
        "Сравни первый квартал 2026 и второй квартал 2025 "
        "по страховым взносам",
        [],
        llm_plan,
    )
    assert corrected.periods == [
        date(2026, 3, 31),
        date(2025, 6, 30),
    ]


def test_llm_retry_usage_and_circuit_breaker_fallback() -> None:
    provider = OpenAICompatibleProvider(
        name="test",
        base_url="https://example.invalid/v1",
        api_key="synthetic-test-key",
        model="synthetic-model",
        timeout=1,
        max_retries=2,
        data_policy="metadata_only",
    )

    class FakeCompletions:
        def __init__(self) -> None:
            self.calls = 0
            self.always_fail = False

        async def create(self, **_kwargs):
            self.calls += 1
            if self.always_fail or self.calls < 2:
                raise TimeoutError("synthetic timeout")
            return SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(content="синтетический ответ")
                    )
                ],
                usage=SimpleNamespace(
                    prompt_tokens=3,
                    completion_tokens=2,
                    total_tokens=5,
                ),
            )

    fake = FakeCompletions()
    provider._client = SimpleNamespace(
        chat=SimpleNamespace(completions=fake)
    )
    response = asyncio.run(
        provider.complete([{"role": "user", "content": "тест"}])
    )
    assert response == "синтетический ответ"
    assert fake.calls == 2
    assert provider.last_usage["total_tokens"] == 5

    fake.always_fail = True
    fallback = asyncio.run(
        provider.complete([{"role": "user", "content": "резерв"}])
    )
    assert fallback == "резерв"
    assert provider.last_fallback is True
    assert provider._breaker.is_open is True

    try:
        asyncio.run(
            provider.complete(
                [{"role": "user", "content": "без fallback"}],
                allow_fallback=False,
            )
        )
    except LLMProviderError as exc:
        assert exc.code == "circuit_open"
    else:
        raise AssertionError("Открытый circuit breaker должен вернуть код ошибки")


def test_gnivc_qwen_uses_proven_no_think_request_contract() -> None:
    settings = Settings(
        llm_provider="gnivc_qwen",
        gnivc_api_key="synthetic-test-key",
        gnivc_query_model="qwen3.6-35b",
        gnivc_narrative_model="qwen3.5-122b",
    )
    provider = GnivcQwenProvider(settings)
    assert provider._client._client.trust_env is False

    class FakeCompletions:
        kwargs: dict

        async def create(self, **kwargs):
            self.kwargs = kwargs
            return SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(content='{"ok": true}')
                    )
                ],
                usage=SimpleNamespace(
                    prompt_tokens=10,
                    completion_tokens=4,
                    total_tokens=14,
                ),
            )

    fake = FakeCompletions()
    provider._client = SimpleNamespace(
        chat=SimpleNamespace(completions=fake)
    )
    original_messages = [
        {"role": "user", "content": "Верни JSON с планом"}
    ]
    response = asyncio.run(
        provider.complete(
            original_messages,
            response_format={"type": "json_object"},
            temperature=0,
            max_tokens=500,
            extra_body={"thinking": {"type": "enabled"}},
            allow_fallback=False,
        )
    )

    assert response == '{"ok": true}'
    assert original_messages[-1]["content"] == "Верни JSON с планом"
    assert fake.kwargs["model"] == "qwen3.6-35b"
    assert fake.kwargs["temperature"] == 0.2
    assert fake.kwargs["top_p"] == 1
    assert fake.kwargs["max_tokens"] == 500
    assert "response_format" not in fake.kwargs
    assert fake.kwargs["extra_body"] == {
        "chat_template_kwargs": {"enable_thinking": False}
    }
    assert fake.kwargs["messages"][-1]["content"].endswith("/no_think")


def test_gnivc_healthcheck_and_narrative_model_are_provider_specific() -> None:
    settings = Settings(
        llm_provider="gnivc_qwen",
        gnivc_api_key="synthetic-test-key",
        gnivc_query_model="qwen3.6-35b",
        gnivc_narrative_model="qwen3.5-122b",
    )
    provider = GnivcQwenProvider(settings)

    class FakeCompletions:
        kwargs: dict

        async def create(self, **kwargs):
            self.kwargs = kwargs
            return SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(content='{"ok": true}')
                    )
                ],
                usage=None,
            )

    fake = FakeCompletions()
    provider._client = SimpleNamespace(
        chat=SimpleNamespace(completions=fake)
    )
    status = asyncio.run(provider.healthcheck())
    narrative = configured_narrative_provider(settings)

    assert status["status"] == "available"
    assert fake.kwargs["max_tokens"] == 1024
    assert fake.kwargs["temperature"] == 0.2
    assert fake.kwargs["extra_body"] == {
        "chat_template_kwargs": {"enable_thinking": False}
    }
    assert isinstance(narrative, GnivcQwenProvider)
    assert narrative.model == "qwen3.5-122b"


def test_gnivc_gigachat_report_editor_uses_local_openai_contract() -> None:
    settings = Settings(
        report_synthesis_provider="gnivc_gigachat",
        report_synthesis_model="GigaChat-Max",
        gnivc_base_url="https://api.ai.gnivc.ru/v1",
        gnivc_api_key="synthetic-gnivc-key",
    )
    provider = GnivcGigaChatProvider(settings)
    assert provider._client._client.trust_env is False

    class FakeCompletions:
        kwargs: dict

        async def create(self, **kwargs):
            self.kwargs = kwargs
            return SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(
                            content='{"summary":"Проверено"}'
                        )
                    )
                ],
                usage=SimpleNamespace(
                    prompt_tokens=12,
                    completion_tokens=4,
                    total_tokens=16,
                ),
            )

    fake = FakeCompletions()
    provider._client = SimpleNamespace(
        chat=SimpleNamespace(completions=fake)
    )
    response_format = {
        "type": "json_schema",
        "schema": {
            "type": "object",
            "properties": {"summary": {"type": "string"}},
            "required": ["summary"],
        },
        "strict": True,
    }
    result = asyncio.run(
        provider.complete(
            [{"role": "user", "content": "Собери отчёт"}],
            response_format=response_format,
            temperature=0.2,
            max_tokens=3000,
            extra_body={"thinking": {"type": "enabled"}},
            allow_fallback=False,
        )
    )

    assert result == '{"summary":"Проверено"}'
    assert fake.kwargs["model"] == "GigaChat-Max"
    assert fake.kwargs["response_format"] == response_format
    assert "extra_body" not in fake.kwargs
    assert fake.kwargs["messages"][-1]["content"] == "Собери отчёт"
    assert "/no_think" not in fake.kwargs["messages"][-1]["content"]
    assert provider.last_usage["total_tokens"] == 16
    assert isinstance(
        configured_report_synthesis_provider(settings),
        GnivcGigaChatProvider,
    )


@REQUIRES_AUTOMSP
def test_automsp_control_pair_parity() -> None:
    data = _report_data()
    assert len(data["table1"]) == 13
    assert len(data["table2"]) == 28
    assert len(data["industries"]) == 20

    rf_total = _row(data, "rf_total")
    msp_total = _row(data, "msp_total")
    msp_share = _row(data, "msp_share")
    micro_total = _row(data, "micro_total")

    assert rf_total["old"] == Decimal("28825.079046")
    assert rf_total["new"] == Decimal("31175.605903908298")
    assert msp_total["old"] == Decimal("6088.880300")
    assert msp_total["new"] == Decimal("7000.085844")
    assert msp_share["old"].quantize(Decimal("0.000001")) == Decimal(
        "21.123551"
    )
    assert msp_share["new"].quantize(Decimal("0.000001")) == Decimal(
        "22.453728"
    )
    assert micro_total["old"] == Decimal("2670.398594")
    assert micro_total["new"] == Decimal("2989.506241")

    budget = _row(data, "msp_budget")
    assert budget["sources"][0]["sheet"] == "Раздел I"
    assert budget["sources"][0]["cell"] == "G9"
    assert budget["sources"][0]["row_code"] == 1010


@REQUIRES_AUTOMSP
def test_generated_docx_has_no_unresolved_tags(tmp_path: Path) -> None:
    docx_path, json_path, snapshot = generate_report(
        TEMPLATE,
        tmp_path,
        OLD,
        NEW,
        find_receipts_workbook(AUTOMSP),
        CODES,
        old_year=2025,
        new_year=2026,
        months_count=6,
    )
    assert docx_path.is_file()
    assert json_path.is_file()
    assert snapshot["report_type"] == "1-msp"

    document = Document(docx_path)
    all_text = "\n".join(
        [
            *(paragraph.text for paragraph in document.paragraphs),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )
    assert "{" not in all_text
    assert len(document.tables) == 4
    assert [len(table.rows) for table in document.tables] == [6, 14, 29, 63]
    assert len(document.tables[2].columns) == 5
    assert all(len(row.cells) == 5 for row in document.tables[2].rows)
    assert not any(
        "tab2_" in cell.text
        for row in document.tables[2].rows
        for cell in row.cells
    )
    appendices = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "Приложения"
    )
    assert appendices.paragraph_format.page_break_before is True
