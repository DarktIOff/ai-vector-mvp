from __future__ import annotations

import asyncio
import copy
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from openpyxl.utils import column_index_from_string
from python_calamine import CalamineWorkbook

from backend.app.config import Settings, get_settings
from backend.app.calculations import ru_number
from backend.app.database import SessionLocal
from backend.app.llm import DeterministicProvider, LLMProviderError
from backend.app.regional import read_regional_report_values
from backend.app.report_service import (
    build_region_vs_rf_report_data,
    interpret_report_command,
    report_available_parameters,
    resolve_report_territory,
    validate_region_report_parameters,
)
from backend.app.reporting import generate_report_from_data


ROOT = Path(__file__).resolve().parents[2]
REQUIRES_HISTORY = pytest.mark.skipif(
    not (ROOT / "rd" / "1_msp_db").is_dir(),
    reason="Для oracle-теста нужен локальный каталог rd/1_msp_db",
)


@pytest.fixture(scope="module")
def yakutia_report() -> dict:
    with SessionLocal() as session:
        return build_region_vs_rf_report_data(
            session,
            get_settings().msp_codes_file,
            territory="Якутия",
            old_period_end=date(2025, 3, 31),
            new_period_end=date(2026, 3, 31),
        )


@REQUIRES_HISTORY
def test_report_catalog_recommends_latest_common_yakutia_pair() -> None:
    with SessionLocal() as session:
        catalog = report_available_parameters(session)
    assert len(catalog["territories"]) == 89
    assert catalog["recommended_pair"] == {
        "old_period_end": "2025-03-31",
        "new_period_end": "2026-03-31",
        "old_label": "январь–март 2025 года",
        "new_label": "январь–март 2026 года",
    }
    yakutia = next(
        item
        for item in catalog["territories"]
        if item["name"] == "Республика Саха (Якутия)"
    )
    assert {"Якутия", "Саха"} <= set(yakutia["aliases"])


@REQUIRES_HISTORY
def test_yakutia_march_2025_2026_oracle(yakutia_report: dict) -> None:
    rows = {item["code"]: item for item in yakutia_report["table1"]}
    expected = {
        "rf_msp_total": ("3126.701142", "3436.389506"),
        "rf_msp_budget": ("2350.660592", "2613.89293"),
        "rf_msp_insurance": ("776.04055", "822.496576"),
        "region_msp_total": ("11.858378", "15.480113"),
        "region_msp_budget": ("8.316153", "11.921862"),
        "region_msp_insurance": ("3.542225", "3.558251"),
        "micro_total": ("5.295155", "7.095227"),
        "small_medium_total": ("6.563223", "8.384886"),
    }
    for code, (old, new) in expected.items():
        assert rows[code]["old"] == Decimal(old)
        assert rows[code]["new"] == Decimal(new)
    assert rows["region_share_rf_msp"]["old"].quantize(
        Decimal("0.000000001")
    ) == Decimal("0.379261639")
    assert rows["region_share_rf_msp"]["new"].quantize(
        Decimal("0.000000001")
    ) == Decimal("0.450476088")

    taxes = {
        item["code"]: (item["old"], item["new"])
        for item in yakutia_report["table2"]
        if item["group"] == "all_msp"
    }
    assert taxes == {
        "region_msp_budget": (Decimal("8.316153"), Decimal("11.921862")),
        "receipts_profit_tax_total": (
            Decimal("1.950909"),
            Decimal("1.822336"),
        ),
        "receipts_ndfl_total": (Decimal("1.765808"), Decimal("2.131902")),
        "receipts_vat_total": (Decimal("2.750292"), Decimal("6.340241")),
        "receipts_special_regimes_total": (
            Decimal("1.382526"),
            Decimal("1.20928"),
        ),
        "receipts_usn": (Decimal("1.19038"), Decimal("0.961332")),
        "receipts_autousn": (Decimal("0.000717"), Decimal("0.19832")),
        "receipts_eshn": (Decimal("0.007817"), Decimal("0.010609")),
        "receipts_psn": (Decimal("0.173198"), Decimal("0.026794")),
        "receipts_npd": (Decimal("0.010345"), Decimal("0.012215")),
    }

    industries = {
        item["base_row_code"]: (
            item["groups"]["all_msp"]["old"],
            item["groups"]["all_msp"]["new"],
        )
        for item in yakutia_report["industries"]
    }
    assert industries == {
        1015: (Decimal("0.061371"), Decimal("0.067391")),
        1036: (Decimal("-0.597498"), Decimal("0.528398")),
        1087: (Decimal("0.62193"), Decimal("0.925706")),
        1255: (Decimal("0.359687"), Decimal("0.151349")),
        1259: (Decimal("0.027136"), Decimal("0.065334")),
        1270: (Decimal("2.592397"), Decimal("3.355875")),
        1295: (Decimal("1.94432"), Decimal("3.117911")),
        1320: (Decimal("1.441425"), Decimal("1.300153")),
        1350: (Decimal("0.190923"), Decimal("0.321565")),
        1363: (Decimal("0.136131"), Decimal("0.116429")),
        1382: (Decimal("0.084866"), Decimal("0.095512")),
        1400: (Decimal("0.35884"), Decimal("0.432347")),
        1405: (Decimal("0.76927"), Decimal("0.98415")),
        1430: (Decimal("0.175334"), Decimal("0.233183")),
        1460: (Decimal("0.014283"), Decimal("0.026411")),
        1465: (Decimal("0.017717"), Decimal("0.028225")),
        1470: (Decimal("0.087718"), Decimal("0.12349")),
        1490: (Decimal("0.014386"), Decimal("0.016839")),
        1497: (Decimal("0.015212"), Decimal("0.028481")),
        1506: (Decimal("0.000434"), Decimal("0.000317")),
    }


def test_small_report_values_remain_visible_without_negative_zero(
    yakutia_report: dict,
    tmp_path: Path,
) -> None:
    data = copy.deepcopy(yakutia_report)
    eshn = next(
        item
        for item in data["table2"]
        if item["code"] == "receipts_eshn" and item["group"] == "all_msp"
    )
    eshn.update(
        old=Decimal("0.024461"),
        new=Decimal("0.023584"),
        difference=Decimal("-0.000877"),
        growth_rate=Decimal("96.414700952536691"),
    )
    share = next(
        item
        for item in data["table1"]
        if item["code"] == "region_share_rf_msp"
    )
    share.update(
        old=Decimal("0.070545693362592567"),
        new=Decimal("0.077245870858505642"),
        difference=Decimal("0.006700177495913075"),
    )

    docx_path, _, _ = generate_report_from_data(
        get_settings().msp_report_template,
        tmp_path,
        data,
        old_year=2025,
        new_year=2026,
        months_count=3,
    )
    document = Document(docx_path)
    summary_labels = [
        row.cells[0].text.strip() for row in document.tables[0].rows[1:]
    ]
    assert [cell.text.strip() for cell in document.tables[0].rows[0].cells] == [
        "ПОКАЗАТЕЛЬ",
        "ЯНВ–МАРТ 2025",
        "ЯНВ–МАРТ 2026",
        "Темп 2026 к 2025, %",
        "Изменение",
    ]
    assert summary_labels == [
        (
            "Российская Федерация — всего от МСП "
            "(консолидированный бюджет + страховые взносы)"
        ),
        (
            "Республика Саха (Якутия) — всего от МСП "
            "(консолидированный бюджет + страховые взносы)"
        ),
        (
            "Республика Саха (Якутия) — поступления в "
            "консолидированный бюджет (без страховых взносов)"
        ),
        "Республика Саха (Якутия) — страховые взносы",
        (
            "Доля региона «Республика Саха (Якутия)» в совокупных "
            "поступлениях от МСП по РФ (бюджет + взносы), %"
        ),
    ]
    assert [
        [cell.text.strip() for cell in row.cells[1:]]
        for row in document.tables[0].rows[2:5]
    ] == [
        ["11,858", "15,480", "130,5", "+3,622"],
        ["8,316", "11,922", "143,4", "+3,606"],
        ["3,542", "3,558", "100,5", "+0,016"],
    ]
    appendix_labels = [
        row.cells[0].text.strip() for row in document.tables[1].rows[1:]
    ]
    assert appendix_labels[1] == (
        "РФ — поступления в консолидированный бюджет "
        "(без страховых взносов)"
    )
    assert appendix_labels[5] == (
        "Республика Саха (Якутия) — поступления "
        "в консолидированный бюджет (без страховых взносов)"
    )
    assert appendix_labels[8] == (
        "Микро — поступления в консолидированный бюджет региона "
        "(без страховых взносов)"
    )
    assert appendix_labels[11] == (
        "Малые и средние без микро — поступления в консолидированный "
        "бюджет региона (без страховых взносов)"
    )
    assert all(
        label != "в консолидированный бюджет РФ"
        for label in appendix_labels
    )
    assert document.paragraphs[2].text.startswith(
        "ТАБЛИЦА: СРАВНЕНИЕ СОВОКУПНЫХ ПОСТУПЛЕНИЙ ОТ МСП"
    )
    assert "КОНСОЛИДИРОВАННЫЙ БЮДЖЕТ + СТРАХОВЫЕ ВЗНОСЫ" in (
        document.paragraphs[2].text
    )
    assert "БЕЗ СТРАХОВЫХ ВЗНОСОВ" in document.paragraphs[17].text
    assert "БЕЗ СТРАХОВЫХ ВЗНОСОВ" in document.paragraphs[18].text
    assert "Темп роста региона был выше темпа по РФ на 20,6 п.п." in (
        document.paragraphs[1].text
    )
    assert "11,922 млрд рублей в консолидированный бюджет плюс" in (
        document.paragraphs[1].text
    )
    assert "3,558 млрд рублей страховых взносов" in (
        document.paragraphs[1].text
    )
    assert "8,316 + 3,542 = 11,858 млрд рублей" in (
        document.paragraphs[1].text
    )
    assert "только поступления в консолидированный бюджет" in (
        document.paragraphs[1].text
    )
    assert [
        cell.text.strip() for cell in document.tables[2].rows[1].cells[1:]
    ] == ["8,316", "11,922", "143,4", "3,606"]
    assert [
        cell.text.strip() for cell in document.tables[3].rows[2].cells[2:6]
    ] == ["8,316", "11,922", "3,606", "143,4"]
    tax_row = next(
        row
        for row in document.tables[2].rows
        if row.cells[0].text.strip() == "ЕСХН"
    )
    assert [cell.text.strip() for cell in tax_row.cells[1:]] == [
        "0,024461",
        "0,023584",
        "96,4",
        "-0,000877",
    ]
    share_row = next(
        row
        for table in document.tables[:2]
        for row in table.rows
        if "Доля региона" in row.cells[0].text
    )
    assert [cell.text.strip() for cell in share_row.cells[1:]] == [
        "0,071%",
        "0,077%",
        "х",
        "+0,007 п.п.",
    ]
    cell_texts = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    assert not any(
        re.search(r"(?<!\d)-0,0+(?!\d)", text) for text in cell_texts
    )
    assert ru_number(Decimal("-0.04")) == "0,0"


def _all_evidence(report: dict) -> list[dict]:
    result: list[dict] = []

    def add(source: dict) -> None:
        result.append(source)
        for child in source.get("lineage", []):
            add(child)

    for row in [*report["table1"], *report["table2"]]:
        for source in row["sources"]:
            add(source)
    for industry in report["industries"]:
        for group in industry["groups"].values():
            for source in group["sources"]:
                add(source)
    return result


@REQUIRES_HISTORY
def test_every_regional_report_evidence_matches_source_cell(
    yakutia_report: dict,
) -> None:
    evidence = _all_evidence(yakutia_report)
    derived = [item for item in evidence if item.get("section") == "III"]
    assert derived
    for item in derived:
        assert item["formula"] == (
            "малые и средние без микро = Раздел I − Раздел II"
        )
        left, right = item["lineage"]
        assert Decimal(item["raw_value"]) == (
            Decimal(left["raw_value"]) - Decimal(right["raw_value"])
        )

    raw = {
        (
            item["absolute_path"],
            item["sheet"],
            item["cell"],
            item["raw_value"],
        ): item
        for item in evidence
        if item.get("cell") not in {None, "—"}
    }
    assert raw
    books: dict[str, CalamineWorkbook] = {}
    sheets: dict[tuple[str, str], list[list[object]]] = {}
    try:
        for (path, sheet, cell, raw_value), item in raw.items():
            assert item.get("file")
            assert item.get("report_column") is not None
            match = re.fullmatch(r"([A-Z]+)(\d+)", cell)
            assert match is not None
            column = column_index_from_string(match.group(1)) - 1
            row = int(match.group(2)) - 1
            if path not in books:
                books[path] = CalamineWorkbook.from_path(path)
            key = (path, sheet)
            if key not in sheets:
                sheets[key] = books[path].get_sheet_by_name(sheet).to_python()
            actual = sheets[key][row][column]
            assert Decimal(str(actual)) == Decimal(raw_value)
    finally:
        for book in books.values():
            book.close()


@pytest.mark.parametrize(
    ("reference", "canonical"),
    [
        ("Якутия", "Республика Саха (Якутия)"),
        ("Москва", "г. Москва"),
        (
            "Карачаево-Черкесская Республика",
            "Карачаево-Черкесская Республика",
        ),
    ],
)
def test_report_territory_aliases(reference: str, canonical: str) -> None:
    with SessionLocal() as session:
        assert resolve_report_territory(session, reference) == canonical


def test_regional_report_rejects_unknown_two_regions_and_bad_periods() -> None:
    with SessionLocal() as session:
        with pytest.raises(ValueError, match="Неизвестный регион"):
            resolve_report_territory(session, "Несуществующая область")
        with pytest.raises(ValueError, match="один регион"):
            resolve_report_territory(session, "Москва и Республика Саха")
        with pytest.raises(ValueError, match="одинаковый месяц"):
            validate_region_report_parameters(
                session,
                territory="Якутия",
                old_period_end=date(2025, 3, 31),
                new_period_end=date(2026, 6, 30),
            )
        with pytest.raises(ValueError, match="Период отсутствует"):
            validate_region_report_parameters(
                session,
                territory="Якутия",
                old_period_end=date(2020, 3, 31),
                new_period_end=date(2021, 3, 31),
            )
        with pytest.raises(ValueError, match="Раздела II"):
            read_regional_report_values(
                session,
                period=date(2022, 12, 31),
                section="II",
                row_codes=[1010],
                report_columns=[4],
                territory="Якутия",
            )


class _FakeQwen:
    name = "gnivc_qwen"

    def __init__(self, response: dict):
        self.response = response
        self.messages: list[dict[str, str]] = []

    async def complete(self, messages, **_kwargs):
        self.messages = messages
        return json_dumps(self.response)


def json_dumps(value: dict) -> str:
    import json

    return json.dumps(value, ensure_ascii=False)


def test_model_first_report_command_uses_only_catalog_and_recommended_pair(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fake = _FakeQwen(
        {
            "status": "ready",
            "report_mode": "region_vs_rf",
            "territory": "Республика Саха (Якутия)",
            "benchmark": "Российская Федерация",
            "old_period_end": "2025-03-31",
            "new_period_end": "2026-03-31",
            "clarification_question": None,
            "no_data_reason": None,
        }
    )
    monkeypatch.setattr(
        "backend.app.report_service.configured_provider", lambda _settings: fake
    )
    with SessionLocal() as session:
        result = asyncio.run(
            interpret_report_command(
                session,
                Settings(llm_provider="gnivc_qwen"),
                "Собери стандартную справку по 1-МСП по Якутии",
            )
        )
    assert result.status == "ready"
    assert result.territory == "Республика Саха (Якутия)"
    assert result.old_period_end == date(2025, 3, 31)
    assert result.new_period_end == date(2026, 3, 31)
    system = fake.messages[0]["content"]
    assert "territories" in system and "comparable_pairs" in system
    assert ".xlsx" not in system
    assert "raw_value" not in system


def test_report_command_has_no_parameter_guessing_fallback(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "backend.app.report_service.configured_provider",
        lambda _settings: DeterministicProvider(),
    )
    with SessionLocal() as session:
        with pytest.raises(LLMProviderError, match="заполните форму вручную"):
            asyncio.run(
                interpret_report_command(
                    session,
                    Settings(llm_provider="deterministic"),
                    "Собери справку по Якутии",
                )
            )
