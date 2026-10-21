from __future__ import annotations

import asyncio
import json
import uuid
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.config import Settings
from backend.app.database import Base, SessionLocal, get_db, init_database
from backend.app.main import app
from backend.app.main import _immutable_docx_response
from backend.app.models import (
    AnalysisProject,
    AnalysisProjectBlock,
    Conversation,
    ConversationTurn,
    Scenario,
)
from backend.app.synthesis_service import synthesize_project
from backend.app.workspace_service import (
    add_project_block,
    answer_clarification,
    build_project_document,
    create_conversation_turn,
    format_report_text_numbers,
    project_section_table,
    project_series_chart,
    scenario_steps,
    write_project_docx,
)


def test_report_prose_rounds_money_scales_and_long_decimals() -> None:
    text = format_report_text_numbers(
        "Поступления 417846632.0000 тыс. руб., "
        "изменение +12500 тыс. руб., рост 32.83968%."
    )
    assert "417,85 млрд руб." in text
    assert "+12,50 млн руб." in text
    assert "32,84 %." in text
    assert "417846632.0000" not in text
    assert "32.83968" not in text
    assert "руб.." not in format_report_text_numbers(
        "2026-06-30 — 5311757712.000000 тыс. руб.."
    )


def test_docx_download_uses_an_immutable_per_request_snapshot(tmp_path) -> None:
    source = tmp_path / "analytical_project.docx"
    original = b"first-docx-version"
    source.write_bytes(original)

    response = _immutable_docx_response(source)
    snapshot = Path(response.path)
    assert snapshot != source
    assert snapshot.read_bytes() == original

    # A concurrent rebuild may atomically replace the canonical project file,
    # but the in-flight response must keep reading its original byte sequence.
    source.write_bytes(b"a-new-and-longer-docx-version")
    assert snapshot.read_bytes() == original

    assert response.background is not None
    asyncio.run(response.background())
    assert not snapshot.exists()


def test_regional_industry_project_table_does_not_repeat_industry() -> None:
    table = project_section_table(
        {
            "kind": "regional_industry_growth_worst",
            "title": "Худшая региональная динамика — Строительство",
            "rows": [
                {
                    "entity": "Строительство — Москва",
                    "industry": "Строительство",
                    "region": "Москва",
                    "growth_rate": "-12.5",
                }
            ],
        }
    )
    assert table[0][0] == "Регион"
    assert table[1][0] == "Москва"
    assert all("Строительство" not in cell for row in table for cell in row)


def test_report_chart_uses_compact_year_axis_for_comparable_periods() -> None:
    chart = project_series_chart(
        "trend",
        [
            {
                "key": "construction",
                "label": (
                    "Поступления · Московская область · все МСП · "
                    "Строительство"
                ),
                "unit": "тыс. руб.",
                "points": [
                    {
                        "period_end": f"{year}-06-30",
                        "label": f"январь–июнь {year}",
                        "value": str(year * 1000),
                    }
                    for year in range(2022, 2027)
                ],
            }
        ],
    )
    assert chart is not None
    assert chart["chart_type"] == "line"
    assert chart["period_context"] == "январь–июнь"
    assert [
        point["axis_label"] for point in chart["series"][0]["points"]
    ] == ["2022", "2023", "2024", "2025", "2026"]


def test_report_chart_uses_grouped_bars_for_one_period() -> None:
    common = "Поступления · Московская область · все МСП · "
    chart = project_series_chart(
        "industries",
        [
            {
                "key": "construction",
                "label": common + "Строительство",
                "unit": "тыс. руб.",
                "points": [
                    {
                        "period_end": "2026-06-30",
                        "label": "январь–июнь 2026",
                        "value": "1200000",
                    }
                ],
            },
            {
                "key": "trade",
                "label": common + "Торговля",
                "unit": "тыс. руб.",
                "points": [
                    {
                        "period_end": "2026-06-30",
                        "label": "январь–июнь 2026",
                        "value": "1800000",
                    }
                ],
            },
        ],
    )
    assert chart is not None
    assert chart["chart_type"] == "bar"
    assert chart["mode"] == "absolute"
    assert chart["context"] == (
        "Поступления · Московская область · все МСП"
    )
    assert [item["label"] for item in chart["series"]] == [
        "Строительство",
        "Торговля",
    ]


def deterministic_settings() -> Settings:
    return Settings(
        llm_provider="deterministic",
        llm_data_policy="none",
        allow_external_llm_data=False,
    )


def test_conversation_api_create_pin_list_and_archive() -> None:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    testing_session = sessionmaker(
        bind=engine, autoflush=False, expire_on_commit=False
    )

    def override_db():
        session = testing_session()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_db
    try:
        with TestClient(app) as client:
            created = client.post("/api/conversations", json={})
            assert created.status_code == 200
            conversation_id = created.json()["id"]
            pinned = client.patch(
                f"/api/conversations/{conversation_id}",
                json={"pinned": True},
            )
            assert pinned.status_code == 200
            assert pinned.json()["pinned"] is True
            listed = client.get("/api/conversations")
            assert [item["id"] for item in listed.json()] == [
                conversation_id
            ]
            archived = client.delete(
                f"/api/conversations/{conversation_id}"
            )
            assert archived.status_code == 200
            assert client.get("/api/conversations").json() == []
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_conversation_api_permanent_delete_cascades_content() -> None:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    testing_session = sessionmaker(
        bind=engine, autoflush=False, expire_on_commit=False
    )

    def override_db():
        session = testing_session()
        try:
            yield session
        finally:
            session.close()

    conversation_id = str(uuid.uuid4())
    turn_id = str(uuid.uuid4())
    project_id = str(uuid.uuid4())
    block_id = str(uuid.uuid4())
    scenario_id = str(uuid.uuid4())
    with testing_session() as session:
        conversation = Conversation(
            id=conversation_id,
            title="Старый диалог",
        )
        turn = ConversationTurn(
            id=turn_id,
            conversation=conversation,
            sequence=1,
            content="Старый вопрос",
        )
        project = AnalysisProject(
            id=project_id,
            conversation=conversation,
            title="Старый проект",
        )
        block = AnalysisProjectBlock(
            id=block_id,
            project=project,
            turn_id=turn_id,
            snapshot_json={},
        )
        scenario = Scenario(
            id=scenario_id,
            title="Сохранённый сценарий",
            source_conversation_id=conversation_id,
        )
        session.add_all(
            [conversation, turn, project, block, scenario]
        )
        session.commit()

    app.dependency_overrides[get_db] = override_db
    try:
        with TestClient(app) as client:
            deleted = client.delete(
                f"/api/conversations/{conversation_id}/permanent"
            )
            assert deleted.status_code == 200
            assert deleted.json() == {
                "id": conversation_id,
                "deleted": True,
            }
            assert (
                client.get(
                    f"/api/conversations/{conversation_id}"
                ).status_code
                == 404
            )
        with testing_session() as session:
            assert session.get(ConversationTurn, turn_id) is None
            assert session.get(AnalysisProject, project_id) is None
            assert session.get(AnalysisProjectBlock, block_id) is None
            scenario = session.get(Scenario, scenario_id)
            assert scenario is not None
            assert scenario.source_conversation_id is None
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_conversation_persists_trend_clarification_and_project() -> None:
    init_database()
    conversation_id = str(uuid.uuid4())
    settings = deterministic_settings()
    with SessionLocal() as session:
        conversation = Conversation(
            id=conversation_id,
            title="Новый диалог",
            pinned=True,
        )
        session.add(conversation)
        session.commit()
        trend = asyncio.run(
            create_conversation_turn(
                session,
                settings,
                conversation,
                "как изменились поступления по мсп за последние 4 года",
            )
        )
        assert trend["status"] == "ready"
        assert trend["query_plan"]["intent"] == "trend"
        assert len(trend["series"][0]["points"]) == 5

        clarification_conversation = Conversation(
            id=str(uuid.uuid4()), title="Новый диалог"
        )
        session.add(clarification_conversation)
        session.commit()
        pending = asyncio.run(
            create_conversation_turn(
                session,
                settings,
                clarification_conversation,
                "Покажи значение",
            )
        )
        assert pending["status"] == "needs_clarification"
        option = next(
            item
            for item in pending["clarification_options"]
            if item["id"] == "budget_receipts"
        )
        pending_turn = session.get(ConversationTurn, pending["turn_id"])
        assert pending_turn is not None
        clarified = asyncio.run(
            answer_clarification(
                session,
                settings,
                clarification_conversation,
                pending_turn,
                option_id=option["id"],
                text=None,
            )
        )
        assert clarified["status"] == "ready"
        assert clarified["query_plan"]["metric_codes"] == [
            "receipts_budget_total"
        ]

        trend_turn = session.get(ConversationTurn, trend["turn_id"])
        assert trend_turn is not None
        project_result = add_project_block(
            session, conversation, trend_turn
        )
        assert project_result["added"] is True
        assert project_result["project"]["blocks_count"] == 1
        assert len(scenario_steps(conversation)) == 1

        session.expunge_all()
        restored = session.get(Conversation, conversation_id)
        assert restored is not None
        assert restored.pinned is True
        assert len(restored.turns) == 1
        assert restored.project is not None

        session.delete(restored)
        clarification_restored = session.get(
            Conversation, clarification_conversation.id
        )
        assert clarification_restored is not None
        session.delete(clarification_restored)
        session.commit()


def test_four_turn_contextual_analytics_chain() -> None:
    init_database()
    settings = deterministic_settings()
    conversation_id = str(uuid.uuid4())
    with SessionLocal() as session:
        conversation = Conversation(
            id=conversation_id, title="Новый диалог"
        )
        session.add(conversation)
        session.commit()
        questions = [
            "по каким отраслям наилучшая ситуация в стране?",
            (
                "какой средний темп роста по всем регионам "
                "по этим отраслям"
            ),
            "как общая ситуация по РФ",
            "В каком регионе лучше всего",
            "Какие регионы выбиваются из общей статистики",
        ]
        results = [
            asyncio.run(
                create_conversation_turn(
                    session, settings, conversation, question
                )
            )
            for question in questions
        ]
        assert [item["status"] for item in results] == ["ready"] * 5
        assert [item["query_plan"]["intent"] for item in results] == [
            "rank",
            "regional_average_growth",
            "overview",
            "rank",
            "outliers",
        ]
        assert results[0]["query_plan"]["analysis_dimension"] == "industry"
        assert {
            item["kind"] for item in results[1]["analysis_sections"]
        } == {"regional_average_growth"}
        assert len(results[1]["analysis_sections"][0]["rows"]) >= 5
        assert all(
            row["regions_count"] > 1
            for row in results[1]["analysis_sections"][0]["rows"]
        )
        assert "средний арифметический прирост" in results[1]["answer"]
        assert {
            item["kind"] for item in results[0]["analysis_sections"]
        } >= {"ranking_value", "ranking_growth"}
        assert results[3]["query_plan"]["analysis_dimension"] == "territory"
        assert {
            item["kind"]
            for item in results[4]["analysis_sections"]
        } == {
            "outliers_value",
            "outliers_growth",
            "outliers_composite",
        }
        session.delete(conversation)
        session.commit()


def test_okved_growth_followup_keeps_receipts_and_full_year_context() -> None:
    init_database()
    settings = deterministic_settings()
    with SessionLocal() as session:
        conversation = Conversation(
            id=str(uuid.uuid4()), title="Новый диалог"
        )
        session.add(conversation)
        session.commit()
        baseline = asyncio.run(
            create_conversation_turn(
                session,
                settings,
                conversation,
                "что по поступлениям мсп за последние 2 полных года?",
            )
        )
        followup = asyncio.run(
            create_conversation_turn(
                session,
                settings,
                conversation,
                "по каким оквэд наилучшая динамика?",
            )
        )

        assert baseline["query_plan"]["intent"] == "compare"
        assert baseline["query_plan"]["periods"] == [
            "2024-12-31",
            "2025-12-31",
        ]
        assert followup["status"] == "ready"
        assert followup["query_plan"]["intent"] == "rank_growth"
        assert followup["query_plan"]["analysis_dimension"] == "industry"
        assert followup["query_plan"]["periods"] == ["2025-12-31"]
        assert followup["query_plan"]["metric_codes"] == [
            "receipts_budget_total",
            "receipts_insurance_total",
        ]
        assert {
            item["kind"] for item in followup["analysis_sections"]
        } == {"ranking_growth", "ranking_low_base"}
        session.delete(conversation)
        session.commit()


def test_dynamics_industries_and_regional_leaders_chain() -> None:
    init_database()
    settings = deterministic_settings()
    with SessionLocal() as session:
        conversation = Conversation(
            id=str(uuid.uuid4()), title="Новый диалог"
        )
        session.add(conversation)
        session.commit()
        questions = [
            "Какая динамика по поступлениям МСП по РФ?",
            "по каким отраслям наибольший рост?",
            "а какие регионы лидируют по росту по этим отраслям?",
        ]
        results = [
            asyncio.run(
                create_conversation_turn(
                    session,
                    settings,
                    conversation,
                    question,
                )
            )
            for question in questions
        ]
        assert [item["status"] for item in results] == ["ready"] * 3
        assert [item["query_plan"]["intent"] for item in results] == [
            "trend",
            "rank_growth",
            "regional_industry_rank_growth",
        ]
        assert len(results[0]["series"][0]["points"]) == 5
        selected_industries = {
            row["row_code"]
            for section in results[1]["analysis_sections"]
            if section["kind"] == "ranking_growth"
            for row in section["rows"]
        }
        assert set(results[2]["query_plan"]["row_codes"]) == selected_industries
        assert results[2]["query_plan"]["metric_codes"] == [
            "receipts_budget_total"
        ]
        assert len(results[2]["analysis_sections"]) == len(
            selected_industries
        )
        assert all(
            section["kind"] == "regional_industry_growth"
            and 1 <= len(section["rows"]) <= 5
            for section in results[2]["analysis_sections"]
        )
        assert "Региональные лидеры роста" in results[2]["answer"]
        session.delete(conversation)
        session.commit()


def test_analytical_brief_is_added_to_project_as_five_tables() -> None:
    init_database()
    settings = deterministic_settings()
    with SessionLocal() as session:
        conversation = Conversation(
            id=str(uuid.uuid4()), title="Сводный аналитический обзор"
        )
        session.add(conversation)
        session.commit()
        result = asyncio.run(
            create_conversation_turn(
                session,
                settings,
                conversation,
                (
                    "Покажи динамику по России, выбери отрасли регионы "
                    "с лучшей/худшей динамикой, покажи где лучше всего, "
                    "где худше всего (где — это по регионам)"
                ),
            )
        )
        assert result["status"] == "ready"
        assert result["query_plan"]["intent"] == "analytical_brief"
        session.refresh(conversation)
        project_result = add_project_block(
            session,
            conversation,
            conversation.turns[-1],
        )
        document = project_result["project"]["document"]
        assert document["table_count"] == 5
        assert [
            block["text"]
            for block in document["blocks"]
            if block["type"] == "paragraph"
            and block["style"] == "Heading 3"
        ] == [
            "Отрасли с лучшей динамикой",
            "Отрасли с худшей динамикой",
            "Лучшие регионы в выбранных отраслях",
            "Худшие регионы в выбранных отраслях",
        ]
        session.delete(conversation)
        session.commit()


def test_report_editor_receives_verified_aggregates_and_strict_schema(
    monkeypatch,
) -> None:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    testing_session = sessionmaker(
        bind=engine, autoflush=False, expire_on_commit=False
    )

    class FakeProvider:
        name = "fake"
        model = "synthesis-test"
        last_usage = {"total_tokens": 10}

        def __init__(self) -> None:
            self.kwargs = {}
            self.messages = []

        async def complete(self, messages, **kwargs):
            self.kwargs = kwargs
            self.messages = messages
            return (
                '{"summary":"Общий вывод","findings":'
                '["Рост подтверждён"],"caveats":[]}'
            )

    provider = FakeProvider()
    monkeypatch.setattr(
        "backend.app.synthesis_service.configured_report_synthesis_provider",
        lambda _settings: provider,
    )
    with testing_session() as session:
        conversation = Conversation(
            id=str(uuid.uuid4()), title="Проект без чисел"
        )
        turn = ConversationTurn(
            id=str(uuid.uuid4()),
            conversation=conversation,
            sequence=1,
            role="user",
            turn_type="analysis",
            status="ready",
            content="Сравни ситуацию",
            plan_json={"intent": "overview"},
            result_json={
                "answer": "Рост подтверждён.",
                "analysis_sections": [
                    {
                        "kind": "ranking_growth",
                        "title": "Лидеры по росту",
                        "methodology": "Сопоставимые периоды.",
                        "rows": [
                            {
                                "entity": "Строительство",
                                "current": "1777477172.000000",
                                "previous": "1595763099.000000",
                                "difference": "181714073.000000",
                                "growth_rate": "11.38728380884812",
                                "unit": "тыс. руб.",
                            }
                        ],
                    }
                ],
                "series": [],
                "numbers_verified": True,
            },
        )
        session.add_all([conversation, turn])
        session.commit()
        add_project_block(session, conversation, turn)
        result = asyncio.run(
            synthesize_project(
                session,
                Settings(
                    llm_provider="fake",
                    llm_data_policy="aggregates",
                    allow_external_llm_data=True,
                ),
                conversation.project,
            )
        )
        assert result["mode"] == "editor:fake"
        assert result["status"] == "success"
        assert result["used_ai"] is True
        assert result["model"] == "synthesis-test"
        assert result["error_code"] is None
        assert "extra_body" not in provider.kwargs
        assert provider.kwargs["response_format"]["type"] == "json_schema"
        assert provider.kwargs["response_format"]["strict"] is True
        assert provider.kwargs["temperature"] == 0.2
        assert "не добавляй ни одного числа" in (
            provider.messages[0]["content"].lower()
        )
        assert "11,39 %" in provider.messages[1]["content"]
        assert "11.38728380884812" not in provider.messages[1]["content"]
        assert any(
            block["id"] == "project-executive-summary"
            for block in result["project"]["document"]["blocks"]
        )
        summary = next(
            block
            for block in result["project"]["document"]["blocks"]
            if block["id"] == "project-executive-summary"
        )
        assert summary["style"] == "Normal"
        table = next(
            block
            for block in result["project"]["document"]["blocks"]
            if block["type"] == "table"
        )
        assert "Индекс" not in table["rows"][0]
        assert table["rows"][1][1] == "1\u00a0777,48"
        assert table["rows"][1][4] == "+11,39 %"
    engine.dispose()


def test_synthesis_edits_from_full_conversation_and_exports_chart(
    monkeypatch,
    tmp_path,
) -> None:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    testing_session = sessionmaker(
        bind=engine, autoflush=False, expire_on_commit=False
    )

    class EditorialProvider:
        name = "fake-editor"
        model = "editorial-test"
        last_usage = {"total_tokens": 42}

        def __init__(self) -> None:
            self.payload: dict = {}
            self.calls = 0

        async def complete(self, messages, **_kwargs):
            self.calls += 1
            self.payload = json.loads(messages[1]["content"])
            first, second = self.payload["blocks"]
            response = {
                    "title": (
                        "География и отраслевая динамика поступлений МСП"
                    ),
                    "subtitle": (
                        "Сравнительный обзор общей динамики и региональных "
                        "различий"
                    ),
                    "summary": "Динамика неоднородна по выбранным срезам.",
                    "findings": [
                        first["verified_answer"],
                        second["verified_answer"],
                    ],
                    "caveats": ["Использованы сопоставимые периоды."],
                    "section_order": [
                        second["block_id"],
                        first["block_id"],
                    ],
                    "sections": [
                        {
                            "block_id": first["block_id"],
                            "heading": "Общая динамика поступлений",
                            "narrative": first["verified_answer"],
                        },
                        {
                            "block_id": second["block_id"],
                            "heading": "Региональные различия",
                            "narrative": second["verified_answer"],
                        },
                    ],
                }
            if self.calls == 1:
                response["title"] += " 999"
            return json.dumps(response, ensure_ascii=False)

    provider = EditorialProvider()
    monkeypatch.setattr(
        "backend.app.synthesis_service.configured_report_synthesis_provider",
        lambda _settings: provider,
    )
    with testing_session() as session:
        conversation = Conversation(
            id=str(uuid.uuid4()), title="Поступления МСП"
        )
        first_turn = ConversationTurn(
            id=str(uuid.uuid4()),
            conversation=conversation,
            sequence=1,
            role="user",
            turn_type="analysis",
            status="ready",
            content="Что по поступлениям МСП за последние годы?",
            plan_json={"intent": "trend"},
            result_json={
                "answer": "Поступления выросли.",
                "query_plan": {
                    "territories": ["Российская Федерация"],
                },
                "resolved_metrics": [
                    {
                        "name_ru": (
                            "Поступления в консолидированный бюджет"
                        )
                    }
                ],
                "analysis_sections": [],
                "series": [
                    {
                        "key": "russia-total",
                        "label": (
                            "Поступления в консолидированный бюджет · "
                            "Российская Федерация · все МСП"
                        ),
                        "unit": "тыс. руб.",
                        "points": [
                            {
                                "period_end": "2025-06-30",
                                "label": "январь–июнь 2025",
                                "value": "500000000",
                            },
                            {
                                "period_end": "2026-06-30",
                                "label": "январь–июнь 2026",
                                "value": "600000000",
                            },
                        ],
                    }
                ],
                "numbers_verified": True,
            },
        )
        second_turn = ConversationTurn(
            id=str(uuid.uuid4()),
            conversation=conversation,
            sequence=2,
            role="user",
            turn_type="analysis",
            status="ready",
            content="А по каким регионам ситуация хуже?",
            plan_json={"intent": "rank"},
            result_json={
                "answer": "Выявлены регионы с худшей динамикой.",
                "query_plan": {
                    "territories": ["Российская Федерация"],
                    "analysis_dimension": "territory",
                },
                "resolved_metrics": [],
                "analysis_sections": [],
                "series": [],
                "numbers_verified": True,
            },
        )
        context_turn = ConversationTurn(
            id=str(uuid.uuid4()),
            conversation=conversation,
            sequence=3,
            role="user",
            turn_type="analysis",
            status="ready",
            content="Сопоставь это со строительством",
            plan_json={"intent": "compare"},
            result_json={
                "answer": "Строительство сопоставлено с общей динамикой.",
                "query_plan": {
                    "territories": ["Российская Федерация"],
                    "analysis_dimension": "industry",
                },
                "resolved_metrics": [],
                "analysis_sections": [],
                "series": [],
                "numbers_verified": True,
            },
        )
        session.add_all(
            [conversation, first_turn, second_turn, context_turn]
        )
        session.commit()
        first_added = add_project_block(session, conversation, first_turn)
        second_added = add_project_block(session, conversation, second_turn)
        first_block_id = first_added["block_id"]
        second_block_id = second_added["block_id"]

        result = asyncio.run(
            synthesize_project(
                session,
                Settings(
                    llm_provider="fake",
                    llm_data_policy="aggregates",
                    allow_external_llm_data=True,
                ),
                conversation.project,
            )
        )

        assert result["used_ai"] is True
        assert result["error_code"] is None
        assert provider.calls == 2
        assert result["project"]["title"] == (
            "География и отраслевая динамика поступлений МСП"
        )
        assert [
            item["question"] for item in provider.payload["conversation_context"]
        ] == [
            "Что по поступлениям МСП за последние годы?",
            "А по каким регионам ситуация хуже?",
            "Сопоставь это со строительством",
        ]
        assert [
            item["selected_for_report"]
            for item in provider.payload["conversation_context"]
        ] == [True, True, False]
        document_blocks = result["project"]["document"]["blocks"]
        assert next(
            item["text"]
            for item in document_blocks
            if item["id"] == "project-title"
        ) == result["project"]["title"]
        assert next(
            index
            for index, item in enumerate(document_blocks)
            if item["id"] == f"{second_block_id}-editorial-heading"
        ) < next(
            index
            for index, item in enumerate(document_blocks)
            if item["id"] == f"{first_block_id}-editorial-heading"
        )
        assert any(
            item["type"] == "chart"
            and item["id"] == f"{first_block_id}-chart"
            for item in document_blocks
        )
        assert not any(
            "research-plan" in item["id"] for item in document_blocks
        )

        docx_settings = deterministic_settings()
        docx_settings.app_data_dir = tmp_path
        docx_settings.prepare_directories()
        path = write_project_docx(
            conversation.project,
            docx_settings,
        )
        exported = Document(path)
        assert len(exported.inline_shapes) == 1
        assert exported.core_properties.author == "ИИ-ВЕКТОР"
        assert (
            exported.sections[0].header.paragraphs[0].text
            == "ИИ-ВЕКТОР  /  АНАЛИТИЧЕСКИЙ ОТЧЁТ 1-МСП"
        )
        assert "проверяемые данные" in (
            exported.sections[0].footer.paragraphs[0].text
        )
        assert any(
            paragraph.text == "ГРАФИК"
            for paragraph in exported.paragraphs
        )
        assert any(
            paragraph.text.startswith("ВЕРСИЯ ")
            and "ПРОВЕРЕННЫХ РЕЗУЛЬТАТОВ" in paragraph.text
            for paragraph in exported.paragraphs
        )
        document_xml = exported._element.xml
        assert 'w:fill="EAF2FF"' in document_xml
        assert 'w:fill="ECF8F2"' in document_xml
        assert 'w:fill="FFF7E6"' in document_xml
        assert 'w:pStyle w:val="ListBullet"' in document_xml
    engine.dispose()


def test_project_docx_uses_rounded_tables_and_real_index_only(
    tmp_path,
) -> None:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    testing_session = sessionmaker(
        bind=engine, autoflush=False, expire_on_commit=False
    )
    with testing_session() as session:
        conversation = Conversation(
            id=str(uuid.uuid4()), title="Форматирование"
        )
        turn = ConversationTurn(
            id=str(uuid.uuid4()),
            conversation=conversation,
            sequence=1,
            role="user",
            turn_type="analysis",
            status="ready",
            content="Сравни отрасли",
            plan_json={"intent": "rank"},
            result_json={
                "answer": "Проверенный вывод.",
                "analysis_sections": [
                    {
                        "kind": "ranking_growth",
                        "title": "Рост",
                        "methodology": (
                            "Малая база определяется по предыдущему периоду."
                        ),
                        "rows": [
                            {
                                "entity": "Строительство",
                                "current": "1777477172",
                                "previous": "1595763099",
                                "difference": "181714073",
                                "growth_rate": "11.3872838",
                                "unit": "тыс. руб.",
                            }
                        ],
                    },
                    {
                        "kind": "outliers_composite",
                        "title": "Диагностика",
                        "rows": [
                            {
                                "entity": "Регион А",
                                "score": "2.53678",
                                "flagged": True,
                            }
                        ],
                    },
                    {
                        "kind": "ranking_growth",
                        "title": "Тот же рост повторно",
                        "methodology": "Повтор таблицы не нужен.",
                        "rows": [
                            {
                                "entity": "Строительство",
                                "current": "1777477172",
                                "previous": "1595763099",
                                "difference": "181714073",
                                "growth_rate": "11.3872838",
                                "unit": "тыс. руб.",
                            }
                        ],
                    },
                ],
                "series": [],
                "numbers_verified": True,
            },
        )
        session.add_all([conversation, turn])
        session.commit()
        add_project_block(session, conversation, turn)
        project = conversation.project
        assert project is not None
        model = build_project_document(project)
        tables = [
            block for block in model["blocks"] if block["type"] == "table"
        ]
        assert len(tables) == 2
        assert "Индекс" not in tables[0]["rows"][0]
        assert tables[0]["rows"][1][1] == "1\u00a0777,48"
        assert tables[0]["rows"][1][4] == "+11,39 %"
        assert "Индекс" in tables[1]["rows"][0]
        assert "2,54" in tables[1]["rows"][1]
        assert any(
            block.get("text", "").startswith("Методика:")
            for block in model["blocks"]
            if block["type"] == "paragraph"
        )

        settings = deterministic_settings()
        settings.app_data_dir = tmp_path
        settings.prepare_directories()
        path = write_project_docx(project, settings)
        document = Document(path)
        assert "Индекс" not in [
            cell.text for cell in document.tables[0].rows[0].cells
        ]
        assert document.tables[0].cell(1, 1).text == "1\u00a0777,48"
        assert document.tables[0].cell(1, 4).text == "+11,39 %"
        header_shading = document.tables[0].cell(0, 0)._tc.tcPr.find(
            qn("w:shd")
        )
        assert header_shading is not None
        assert header_shading.get(qn("w:fill")) == "17324D"
        assert document.tables[0].rows[0]._tr.trPr.find(
            qn("w:tblHeader")
        ) is not None
    engine.dispose()
